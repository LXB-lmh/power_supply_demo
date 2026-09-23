# -*- coding: utf-8 -*-
"""Word 表格原样拷贝：颜色、列宽、单元格样式跟源 docx，不单填格子。

与图/公式一样，解析时记 source_index，成文时 deepcopy w:tbl 并重映射关系。
Excel / 合并生成的表无 source_index 时仍走 add_table 填字。
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from chapters.common.drawings import _clone_part, _remap_rids, _rids_in
from chapters.common.table_layout import CONTENT_TWIPS
from parsers.document_model import Block, DocumentModel


def table_flow_item(doc: DocumentModel, block: Block) -> dict[str, Any]:
    """表格进 flow，带上回拷所需路径与 body 下标。"""
    item: dict[str, Any] = {"kind": "table", "rows": [list(r) for r in (block.rows or [])]}
    merges = getattr(block, "vmerge", None)
    if merges is not None:
        item["vmerge"] = [list(m) for m in merges]
    idx = int(block.source_index) if (block.source_index is not None and block.source_index >= 0) else -1
    if idx >= 0 and (doc.source_path or "").lower().endswith((".docx", ".doc")):
        item["source_path"] = doc.source_path or ""
        item["source_index"] = idx
    return item


def _insert_body_element(dest_doc: Document, cloned) -> None:
    placeholder = dest_doc.add_paragraph()
    parent = placeholder._element.getparent()
    idx = list(parent).index(placeholder._element)
    parent.remove(placeholder._element)
    parent.insert(idx, cloned)


def resolve_docx_for_copy(src_path: str | Path) -> Path | None:
    """回拷只认 docx。.doc 先找已转换文件，没有再转一次。"""
    src_path = Path(src_path)
    if src_path.suffix.lower() == ".docx" and src_path.is_file():
        return src_path
    if src_path.suffix.lower() != ".doc":
        return None
    from parsers.office_convert import converted_sidecar, convert_legacy

    hit = converted_sidecar(src_path, None, "docx")
    if hit is not None:
        return hit
    if not src_path.is_file():
        return None
    try:
        return convert_legacy(src_path, src_path.parent / "_converted")
    except Exception:
        return None


def copy_table_from_body_index(src_path: str | Path, source_index: int, dest_doc: Document) -> bool:
    """按源文档 body 子节点下标整表拷贝（含底纹、列宽、边框）。"""
    if source_index < 0:
        return False
    resolved = resolve_docx_for_copy(src_path)
    if resolved is None:
        return False
    try:
        src = Document(str(resolved))
    except Exception:
        return False
    children = list(src.element.body.iterchildren())
    if source_index >= len(children):
        return False
    src_elm = children[source_index]
    if src_elm.tag != qn("w:tbl"):
        return False
    cloned = deepcopy(src_elm)
    cache: dict[int, Any] = {}
    rid_map: dict[str, str] = {}
    for rid in _rids_in(cloned):
        if rid not in src.part.rels:
            continue
        rel = src.part.rels[rid]
        if rel.is_external:
            rid_map[rid] = dest_doc.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            new_part = _clone_part(rel.target_part, dest_doc.part.package, cache)
            rid_map[rid] = dest_doc.part.relate_to(new_part, rel.reltype)
    _remap_rids(cloned, rid_map)
    _insert_body_element(dest_doc, cloned)
    return True


# 2025 触网年报表3-2 实测 grid：状态参数/权重窄、评分规则宽、中间隔列极窄。
WEIGHT_RULE_WIDTHS_7 = [679, 300, 3014, 142, 600, 266, 2960]
# 2025 触网年报表3-1 设备评级（比均分 8294 更窄，与后表视觉同宽量级）
GRADE_RANGE_WIDTHS_2 = [2220, 4416]


def _is_weight_rule_table(rows: list[list[str]] | None) -> bool:
    if not rows:
        return False
    blob = "".join(str(c or "") for c in rows[0])
    return "状态参数" in blob and ("评分计算" in blob or "权重" in blob)


def is_plan_exec_table(rows: list[list[str]] | None) -> bool:
    """生产计划执行表：计划数量 + 完成率/完成数量。"""
    if not rows or len(rows) < 2:
        return False
    head = "".join(str(c or "") for c in (rows[0] or []))
    return "计划数量" in head and ("完成率" in head or "完成数量" in head)


def plan_exec_widths(ncols: int, *, source: list[int] | None = None) -> list[int]:
    """按原材料列宽比例拉到年报正文宽，避免 完成率 被挤成 1000%。"""
    from chapters.common.table_layout import even_widths

    n = max(int(ncols or 0), 1)
    if source and len(source) >= n and sum(int(w) for w in source[:n]) > 0:
        raw = [max(200, int(w)) for w in source[:n]]
        s = sum(raw)
        widths = [max(260, int(CONTENT_TWIPS * w / s)) for w in raw]
        widths[-1] += CONTENT_TWIPS - sum(widths)
        return widths
    if n == 7:
        return [680, 1180, 1520, 1520, 1520, 1040, 834]
    return even_widths(n, CONTENT_TWIPS)


def _table_grid_widths(table) -> list[int]:
    if table is None:
        return []
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is None:
        return []
    out: list[int] = []
    for col in grid.findall(qn("w:gridCol")):
        try:
            out.append(int(col.get(qn("w:w")) or 0))
        except (TypeError, ValueError):
            out.append(0)
    return out


def _mark_repeat_header(table, n: int = 1) -> None:
    """跨页续表重复表头，避免后半段看起来像少了线路。"""
    if table is None:
        return
    for ri, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        for child in list(tr_pr):
            if child.tag == qn("w:tblHeader"):
                tr_pr.remove(child)
        if ri < n:
            el = OxmlElement("w:tblHeader")
            el.set(qn("w:val"), "true")
            tr_pr.append(el)


def apply_plan_exec_layout(table, rows: list[list[str]] | None = None) -> None:
    """计划表跟原材料列比例，固定列宽并重复表头。"""
    if table is None:
        return
    head = list(rows[0]) if rows else [c.text for c in table.rows[0].cells]
    blob = "".join(str(c or "") for c in head)
    if "计划数量" not in blob or ("完成率" not in blob and "完成数量" not in blob):
        return
    ncols = max(len(head), len(table.columns) if table.columns else 0)
    src_grid = _table_grid_widths(table)
    widths = plan_exec_widths(ncols, source=None if ncols == 7 else src_grid)
    apply_table_widths(table, widths, jc="center")
    _mark_repeat_header(table, 1)
    from chapters.common.docx_style import _set_tbl_cell_mar

    _set_tbl_cell_mar(table, left=40, right=40)


def is_cycle_maintain_table(rows: list[list[str]] | None) -> bool:
    """触网/供电维护周期表：大类+工作项目+周期+内容，或三列工作项目表。"""
    if not rows:
        return False
    blob = "".join(str(c or "") for c in rows[0])
    return "维护周期" in blob and "维护内容" in blob and ("工作项目" in blob or "大类" in blob)


def apply_cycle_maintain_layout(table, rows: list[list[str]] | None = None) -> None:
    """4.2.2 跟人工年报：四列 7429、去掉部门稿单元格边距，只留一张。"""
    if table is None:
        return
    from chapters.common.docx_style import _set_tbl_cell_mar
    from chapters.common.table_layout import pad_widths
    from chapters.overhead.style import OH_CYCLE_4COL_WIDTHS, OH_CYCLE_TBL_W

    head = rows[0] if rows else [c.text for c in table.rows[0].cells]
    if not is_cycle_maintain_table([head]):
        return
    ncols = max(len(head), len(table.columns) if table.columns else 0)
    if ncols == 4:
        widths = list(OH_CYCLE_4COL_WIDTHS)
    else:
        widths = pad_widths(list(OH_CYCLE_4COL_WIDTHS), ncols, total=OH_CYCLE_TBL_W)
    apply_table_widths(table, widths, jc="center")
    _set_tbl_cell_mar(table)


def collapse_cycle_maintain_tables(flow: list[dict[str, Any]] | None) -> list[dict[str, Any]]:
    """维护周期节只留一张最完整的表，避免部门稿+补文件叠出两张。"""
    items = [x for x in (flow or []) if isinstance(x, dict)]
    idxs = [i for i, x in enumerate(items) if x.get("kind") == "table" and is_cycle_maintain_table(x.get("rows"))]
    if len(idxs) <= 1:
        return items
    best = max(idxs, key=lambda i: (len(items[i].get("rows") or []), i))
    return [
        x
        for i, x in enumerate(items)
        if i == best or not (x.get("kind") == "table" and is_cycle_maintain_table(x.get("rows")))
    ]


def default_layout_widths(rows: list[list[str]] | None) -> list[int] | None:
    """重建时套去年列宽，避免 Word 均分把规则列拆成竖条。"""
    if not rows:
        return None
    head = [str(c or "").strip() for c in rows[0]]
    if _is_weight_rule_table(rows) and len(head) == 7:
        return list(WEIGHT_RULE_WIDTHS_7)
    if len(head) == 2 and "设备评级" in head[0] and "评级范围" in head[1]:
        return list(GRADE_RANGE_WIDTHS_2)
    if is_cycle_maintain_table(rows) and len(head) == 4:
        from chapters.overhead.style import OH_CYCLE_4COL_WIDTHS

        return list(OH_CYCLE_4COL_WIDTHS)
    if is_plan_exec_table(rows):
        return plan_exec_widths(len(head))
    return None


def _status_header_row_count(rows: list[list[str]]) -> int:
    if len(rows) >= 2:
        r1 = "".join(str(c or "") for c in rows[1])
        if any(k in r1 for k in ("占比", "锚段数", "区段数", "数量")) and "号线" not in r1:
            return 2
    return 1


def _status_header_vmerge(ncols: int, header_rows: int) -> list[tuple[int, int, int]]:
    """双行表头：线路、锚段数量列竖合并。"""
    if header_rows < 2 or ncols < 2:
        return []
    return [(0, 0, 1), (1, 0, 1)]


def _apply_status_header_hmerge(table, ncols: int) -> None:
    """第一行 A/B/C/D 各合并两列，不要显示成 A A B B C C D D。"""
    from chapters.common.docx_style import _hmerge_row

    col = 2
    while col + 1 < ncols:
        _hmerge_row(table, 0, col, col + 1)
        col += 2


def add_table_full_rows(
    dest_doc: Document,
    rows: list[list[str]],
    *,
    widths: list[int] | None = None,
    vmerge: list[tuple[int, int, int]] | None = None,
    font_pt: float | None = None,
    jc: str | None = "center",
    allow_infer_vmerge: bool = True,
) -> None:
    """整表按行写入（含双行表头），避免 add_table 把第二行表头当成数据行。"""
    if not rows:
        return
    from chapters.common.docx_style import BODY_PT, _format_table_cells, _set_table_width

    ncols = max(len(r) for r in rows)
    norm: list[list[str]] = []
    for r in rows:
        cells = [str(c or "").strip() for c in r]
        if len(cells) < ncols:
            cells.extend([""] * (ncols - len(cells)))
        norm.append(cells[:ncols])
    header_rows = _status_header_row_count(norm)
    table = dest_doc.add_table(rows=len(norm), cols=ncols)
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False
    for ri, row in enumerate(norm):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val
    merges = list(vmerge) if vmerge else []
    if not merges and allow_infer_vmerge:
        from chapters.common.word import infer_vmerges

        merges = infer_vmerges(norm, header_rows=header_rows)
    if header_rows >= 2:
        head_m = _status_header_vmerge(ncols, header_rows)
        merges = head_m + [m for m in merges if m not in head_m]
    from chapters.common.docx_style import _vmerge_column

    for col, start, end in merges:
        _vmerge_column(table, col, start, end)
    if header_rows >= 2 and ncols >= 10:
        _apply_status_header_hmerge(table, ncols)
    size = font_pt if font_pt is not None else BODY_PT
    _format_table_cells(table, header_bold=True, font_pt=size)
    if widths:
        _set_table_width(table, [int(w) for w in widths], jc=jc or "center")
    from chapters.common.docx_style import apply_weight_rule_shading

    if _is_weight_rule_table(norm):
        apply_weight_rule_shading(table)


def apply_table_widths(table, widths: list[int] | None, *, jc: str | None = "center") -> None:
    """整表回拷后仍强制 grid/tcW，避免部门材料里极窄列把数字拆行。"""
    if not table or not widths:
        return
    from chapters.common.docx_style import _set_table_width

    _set_table_width(table, [int(w) for w in widths], jc=jc)


def write_table_item(dest_doc: Document, item: dict[str, Any], *, domain_id: str | None = None) -> bool:
    """优先原样拷表；失败或无溯源时按 rows + vmerge 重建。"""
    domain_id = domain_id or (str(item.get("domain_id") or "") or None)
    path = str(item.get("source_path") or "")
    raw_idx = item.get("source_index")
    idx = int(raw_idx) if raw_idx is not None and int(raw_idx) >= 0 else -1
    widths = item.get("widths") if isinstance(item.get("widths"), list) else None
    jc = item.get("jc") if item.get("jc") is not None else "center"
    if path and idx >= 0 and not item.get("force_rebuild"):
        if copy_table_from_body_index(path, idx, dest_doc):
            if dest_doc.tables:
                # 原表已带列宽；只有调用方显式给宽度或维护周期表才改。
                if widths:
                    apply_table_widths(dest_doc.tables[-1], widths, jc=jc)
                apply_cycle_maintain_layout(dest_doc.tables[-1], item.get("rows"))
                apply_plan_exec_layout(dest_doc.tables[-1], item.get("rows"))
            return True
    rows = item.get("rows") or []
    if not rows or len(rows) < 1:
        return False
    if not widths:
        widths = default_layout_widths(rows)
    header_rows = _status_header_row_count(rows)
    allow_infer = bool(
        item.get("force_rebuild") or _is_weight_rule_table(rows) or header_rows >= 2
    )
    if header_rows >= 2 or item.get("force_rebuild") or _is_weight_rule_table(rows):
        add_table_full_rows(
            dest_doc,
            rows,
            widths=widths,
            vmerge=item.get("vmerge") if isinstance(item.get("vmerge"), list) else None,
            font_pt=item.get("font_pt") if item.get("font_pt") is not None else 10.5,
            jc=jc,
            allow_infer_vmerge=allow_infer,
        )
        if dest_doc.tables:
            if is_cycle_maintain_table(rows):
                apply_cycle_maintain_layout(dest_doc.tables[-1], rows)
            apply_plan_exec_layout(dest_doc.tables[-1], rows)
        return True
    from chapters.common.word import add_table

    raw = item.get("vmerge")
    vmerge: list[tuple[int, int, int]] | None = None
    if raw is not None:
        vmerge = []
        for m in raw:
            if isinstance(m, (list, tuple)) and len(m) >= 3:
                vmerge.append((int(m[0]), int(m[1]), int(m[2])))
    elif not allow_infer:
        vmerge = []
    add_table(
        dest_doc,
        rows,
        vmerge=vmerge,
        widths=widths,
        jc=jc,
        font_pt=item.get("font_pt"),
        cell_mar=0 if is_cycle_maintain_table(rows) else None,
        domain_id=domain_id,
    )
    if dest_doc.tables:
        if is_cycle_maintain_table(rows):
            apply_cycle_maintain_layout(dest_doc.tables[-1], rows)
        apply_plan_exec_layout(dest_doc.tables[-1], rows)
    return True
