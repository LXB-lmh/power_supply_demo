# -*- coding: utf-8 -*-
"""年报 Word 共用字体、段落和表格样式。

宋体 12 磅、1.5 倍行距、标题不加粗编号。去掉 theme 字体，避免 WPS 把东亚文字换成等线。
单元格里的「（待填）」黄底，对应措施未填；不要改成正文里的【材料未提供】句。
"""
from __future__ import annotations

from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HEADING_PT = {1: 12, 2: 12, 3: 12, 4: 12}
BODY_PT = 12
BODY_FONT = "宋体"
BLACK = RGBColor(0x00, 0x00, 0x00)
THEME_FONT_ATTRS = (
    qn("w:asciiTheme"),
    qn("w:hAnsiTheme"),
    qn("w:eastAsiaTheme"),
    qn("w:cstheme"),
)
# 表内占位：黄底单元格。成文默认不往正文写这两句；若材料或旧稿带进来，仍按黄标处理。
PENDING_MARKS = {"（待填）", "【材料未提供，本节留空，禁止编造。】"}
# 2025 多数表合计宽度，与 format.CONTENT_TWIPS 一致，不要用页面 100%。
CONTENT_TWIPS = 8294


def _hex_color(color) -> str:
    if color is None:
        return "000000"
    if isinstance(color, RGBColor):
        return f"{int(color[0]):02X}{int(color[1]):02X}{int(color[2]):02X}"
    text = str(color).replace("#", "").replace(",", "")
    return text if len(text) == 6 else "000000"


def _strip_theme_fonts(rFonts) -> None:
    """删掉 theme 字体引用。不删的话 WPS 会把东亚文字显示成等线，对不上 2025 宋体。"""
    if rFonts is None:
        return
    for attr in THEME_FONT_ATTRS:
        if attr in rFonts.attrib:
            del rFonts.attrib[attr]


def _set_rfonts(rPr, name: str) -> None:
    rFonts = rPr.get_or_add_rFonts()
    _strip_theme_fonts(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:hint"), "eastAsia")


def _set_color(rPr, hex_color: str) -> None:
    for child in list(rPr):
        if child.tag == qn("w:color"):
            rPr.remove(child)
    el = OxmlElement("w:color")
    el.set(qn("w:val"), hex_color)
    rPr.append(el)


def _set_sz(rPr, pt: float) -> None:
    half = str(int(pt * 2))
    for tag in ("w:sz", "w:szCs"):
        for child in list(rPr):
            if child.tag == qn(tag):
                rPr.remove(child)
        el = OxmlElement(tag)
        el.set(qn("w:val"), half)
        rPr.append(el)


def _apply_body_paragraph(paragraph) -> None:
    """正文：1.5 倍行距、首行 24 磅（两个宋体字），与 2025 年报一致。"""
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(24)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _apply_heading_paragraph(paragraph) -> None:
    """标题：同样 1.5 倍行距，但无首行缩进，避免「3.3.2 ……」被缩进去。"""
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Cm(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _clear_style_numbering(style) -> None:
    p_pr = style.element.find(qn("w:pPr"))
    if p_pr is None:
        return
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)


def _set_doc_defaults(doc: Document) -> None:
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles_el.insert(0, defaults)
    r_def = defaults.find(qn("w:rPrDefault"))
    if r_def is None:
        r_def = OxmlElement("w:rPrDefault")
        defaults.append(r_def)
    rPr = r_def.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        r_def.append(rPr)
    _set_rfonts(rPr, BODY_FONT)
    _set_sz(rPr, BODY_PT)
    _set_color(rPr, "000000")


def _set_style_font(style, *, name: str = BODY_FONT, size: float = BODY_PT, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    if bold is not None:
        style.font.bold = bold
    rPr = style.element.get_or_add_rPr()
    _set_rfonts(rPr, name)
    _set_sz(rPr, size)
    _set_color(rPr, "000000")


def _set_font(doc: Document) -> None:
    """Normal / Heading 都改成宋体 12、不自动编号。标题不加粗，与 2025 供电年报一致。"""
    _set_doc_defaults(doc)
    normal = doc.styles["Normal"]
    _set_style_font(normal, name=BODY_FONT, size=BODY_PT, bold=False)
    nf = normal.paragraph_format
    nf.line_spacing = 1.5
    nf.space_before = Pt(0)
    nf.space_after = Pt(0)
    nf.first_line_indent = Pt(24)
    try:
        nf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    except Exception:
        pass
    for level, size in HEADING_PT.items():
        heading = doc.styles[f"Heading {level}"]
        _set_style_font(heading, name=BODY_FONT, size=size, bold=False)
        _clear_style_numbering(heading)
        hf = heading.paragraph_format
        hf.line_spacing = 1.5
        hf.space_before = Pt(0)
        hf.space_after = Pt(0)
        hf.first_line_indent = Cm(0)


def _run(paragraph, text: str, *, bold=False, size=12, name="宋体", color=None, highlight=False, vert: str | None = None):
    """写一个 run。highlight=True 才黄底，用于不通顺句或黄标题，不是整节铺黄。"""
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    if vert == "superscript":
        run.font.superscript = True
    elif vert == "subscript":
        run.font.subscript = True
    rPr = run._element.get_or_add_rPr()
    _set_rfonts(rPr, name)
    _set_sz(rPr, size)
    if color is not None:
        run.font.color.rgb = color
        _set_color(rPr, _hex_color(color))
    else:
        _set_color(rPr, "000000")
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def _set_cell_shading(cell, fill: str, *, font_hex: str | None = None, bold: bool | None = None) -> None:
    """单元格底纹。表 3-2 等权重表对照 2025 触网年报：表头 2F5597 白字。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:shd"):
            tc_pr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    tc_pr.append(shd)
    if font_hex is None and bold is None:
        return
    rgb = None
    if font_hex:
        hx = font_hex.replace("#", "")
        rgb = RGBColor(int(hx[0:2], 16), int(hx[2:4], 16), int(hx[4:6], 16))
    for p in cell.paragraphs:
        for run in p.runs:
            if bold is not None:
                run.bold = bold
            if rgb is not None:
                run.font.color.rgb = rgb
                rPr = run._element.get_or_add_rPr()
                _set_color(rPr, font_hex)


def apply_weight_rule_shading(table) -> None:
    """表3-2/3-3/3-4：表头深蓝、奇数行 D2DEEF、偶数行 EAEFF7，空隔列同表头色。"""
    if table is None or not table.rows:
        return
    header = table.rows[0]
    spacer = [i for i, c in enumerate(header.cells) if not (c.text or "").strip()]
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            if ci in spacer or ri == 0:
                _set_cell_shading(cell, "2F5597", font_hex="FFFFFF", bold=True)
            elif ri % 2 == 1:
                _set_cell_shading(cell, "D2DEEF", font_hex="000000", bold=False)
            else:
                _set_cell_shading(cell, "EAEFF7", font_hex="000000", bold=False)


def _set_cell_valign(cell, val: str = "center") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:vAlign"):
            tc_pr.remove(child)
    align = OxmlElement("w:vAlign")
    align.set(qn("w:val"), val)
    tc_pr.append(align)


def _set_tbl_cell_mar(table, *, left: int | None = None, right: int | None = None) -> None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        return
    for child in list(tbl_pr):
        if child.tag == qn("w:tblCellMar"):
            tbl_pr.remove(child)
    if left is None and right is None:
        return
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("left", left), ("right", right)):
        if val is None:
            continue
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(int(val)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _vmerge_column(table, col: int, start: int, end: int) -> None:
    if end <= start:
        return
    for row in range(start + 1, end + 1):
        table.cell(row, col).text = ""
    table.cell(start, col).merge(table.cell(end, col))


def _hmerge_row(table, row: int, start: int, end: int) -> None:
    """同一行左右合并。表 3-5a 第一行 A/B/C/D 各跨「锚段数+占比」两列。"""
    if end <= start:
        return
    keep = (table.cell(row, start).text or "").strip()
    for col in range(start + 1, end + 1):
        table.cell(row, col).text = ""
    table.cell(row, start).merge(table.cell(row, end))
    if keep:
        table.cell(row, start).text = keep


def _align_of(name: str | None):
    key = (name or "center").lower()
    if key in {"both", "justify"}:
        return WD_ALIGN_PARAGRAPH.JUSTIFY
    if key == "left":
        return WD_ALIGN_PARAGRAPH.LEFT
    if key == "right":
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.CENTER


def _even_col_widths(n: int, total: int = CONTENT_TWIPS) -> list[int]:
    if n <= 0:
        return []
    base, rem = divmod(int(total), n)
    return [base + (1 if i < rem else 0) for i in range(n)]


def _set_table_width(table, widths: list[int] | None = None, *, jc: str | None = "center") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    for child in list(tbl_pr):
        if child.tag in {qn("w:tblW"), qn("w:jc"), qn("w:tblLayout")}:
            tbl_pr.remove(child)
    tbl_w = OxmlElement("w:tblW")
    if widths:
        tbl_w.set(qn("w:w"), str(sum(int(w) for w in widths)))
        tbl_w.set(qn("w:type"), "dxa")
    else:
        tbl_w.set(qn("w:w"), "5000")
        tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    if jc:
        el = OxmlElement("w:jc")
        el.set(qn("w:val"), jc)
        tbl_pr.append(el)
    if not widths:
        return
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        return
    cols = grid.findall(qn("w:gridCol"))
    for col, width in zip(cols, widths):
        col.set(qn("w:w"), str(int(width)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            for child in list(tc_pr):
                if child.tag == qn("w:tcW"):
                    tc_pr.remove(child)
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(int(width)))
            tc_w.set(qn("w:type"), "dxa")
            tc_pr.append(tc_w)


def _format_table_cells(
    table,
    *,
    header_bold: bool = False,
    col_align: list[str] | None = None,
    font_pt: float = BODY_PT,
) -> None:
    size = font_pt or BODY_PT
    for r_i, row in enumerate(table.rows):
        for c_i, cell in enumerate(row.cells):
            _set_cell_valign(cell, "center")
            align = _align_of((col_align[c_i] if col_align and c_i < len(col_align) else None) if r_i else "center")
            for p in cell.paragraphs:
                p.alignment = align
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                if not p.runs:
                    if p.text:
                        _run(p, p.text, bold=header_bold and r_i == 0, size=size, name=BODY_FONT)
                    continue
                for run in p.runs:
                    run.bold = header_bold and r_i == 0
                    run.font.size = Pt(size)
                    run.font.name = BODY_FONT
                    rPr = run._element.get_or_add_rPr()
                    _set_rfonts(rPr, BODY_FONT)
                    _set_sz(rPr, size)
                    _set_color(rPr, "000000")


def _table(
    doc: Document,
    headers: list[str],
    rows: list[list[Any]],
    *,
    widths: list[int] | None = None,
    header_bold: bool = False,
    col_align: list[str] | None = None,
    jc: str | None = "center",
    cell_mar: int | None = None,
    vmerge: list[tuple[int, int, int]] | None = None,
    font_pt: float = BODY_PT,
) -> None:
    """固定列宽网格表。单元格等于 PENDING_MARKS 时黄底，对应「待填」而不是缺整节。"""
    if not widths:
        widths = _even_col_widths(len(headers))
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i].cells[c_i]
            text = "" if val is None else str(val)
            cell.text = text
            if text in PENDING_MARKS:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for col, start, end in vmerge or []:
        _vmerge_column(table, col, start, end)
    _format_table_cells(table, header_bold=header_bold, col_align=col_align, font_pt=font_pt)
    _set_table_width(table, widths, jc=jc)
    if cell_mar is not None:
        _set_tbl_cell_mar(table, left=cell_mar, right=cell_mar)
