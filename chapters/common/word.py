# -*- coding: utf-8 -*-
"""年报各章共用的 Word 写法：出处只标在标题后括号；缺材料或不通顺标黄。

黄标三层，不要混：
- 黄标题：整节没有材料，或该节里有不通顺句（小节标题一起黄）；
- 黄句子：只标这一句（话说一半、待提供、缺故障起数、标点坏）；
- 原材料标黄：评估材料该段本身标黄，年报整段黄并加（原材料标黄）。
不要在正文里写【材料未提供】或段末出处。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

from chapters.common.docx_style import (
    BODY_FONT,
    BODY_PT,
    _apply_body_paragraph,
    _apply_heading_paragraph,
    _run,
    _set_font,
)

LEAD_IN = re.compile(
    r"^(具体如下|详情如下|原因|措施|故障现象|故障原因分析|整改措施及建议|分别是)[:：]?$"
)
_NUMERIC_CELL = re.compile(r"^[\d\.％%\s/,，\-—]+$")


# 无材料原 vMerge 时（如 Excel），只推断分类列，勿合物料名称/规格/数量/状态值等同文列。
_MERGEABLE_COL_HEADER = re.compile(r"专业|大类|中类|小类|类型|线路|区段|类别")
_NON_MERGE_COL_HEADER = re.compile(
    r"物料|规格|型号|数量|序号|单位|计量|编号|日期|名称|结果|描述|步骤|措施|备注|状态值|状态等级"
)


def _col_allows_infer_vmerge(header: str) -> bool:
    h = (header or "").strip()
    if not h or not _MERGEABLE_COL_HEADER.search(h):
        return False
    # 「设备大类」可合；「物料名称」虽含「名称」但无大/中/小类词 → 已在上一行挡掉
    if _NON_MERGE_COL_HEADER.search(h) and not re.search(r"大类|中类|小类|专业|类型|线路|区段", h):
        return False
    return True


def _cell_text(rows: list[list[str]], row: int, col: int) -> str:
    if row < 0 or row >= len(rows) or col < 0:
        return ""
    line = rows[row]
    if col >= len(line):
        return ""
    return str(line[col] or "").strip()


def infer_vmerges(
    rows: list[list[str]] | None,
    *,
    header_rows: int = 1,
) -> list[tuple[int, int, int]]:
    """无材料原合并信息时，按同文推断竖合并（兜底）。

    Word 材料应优先用解析出的真实 vMerge；本函数主要服务 Excel 等。
    只合表头像分类的列（专业/大类/中类/线路/区段），避免把「物料名称」等同文列误并。
    右侧分类列只在左侧列也相同的连续段里合并（1号线正线与 2号线正线不相邻合）。
    返回 (列, 起行, 止行)，含表头行下标，与 add_table / _vmerge_column 一致。
    """
    start = max(1, int(header_rows or 1))
    if not rows or len(rows) < start + 2:
        return []
    header = rows[start - 1] if start >= 1 else rows[0]
    ncols = max((len(r) for r in rows), default=0)
    out: list[tuple[int, int, int]] = []
    for col in range(ncols):
        head = str(header[col] if col < len(header) else "")
        if not _col_allows_infer_vmerge(head):
            continue
        i = start
        while i < len(rows):
            val = _cell_text(rows, i, col)
            if not val or _NUMERIC_CELL.match(val):
                i += 1
                continue
            j = i + 1
            while j < len(rows):
                if _cell_text(rows, j, col) != val:
                    break
                left_changed = False
                for lc in range(col):
                    lh = str(header[lc] if lc < len(header) else "")
                    if not _col_allows_infer_vmerge(lh):
                        continue
                    if _cell_text(rows, j, lc) != _cell_text(rows, i, lc):
                        left_changed = True
                        break
                if left_changed:
                    break
                j += 1
            if j - 1 > i:
                out.append((col, i, j - 1))
            i = j
    return out


FAULT_GAP = re.compile(r"故障起(?!\d)")
CUT_END = re.compile(r"[，、——]$|[的和与及或为在将把从对以]$")
PENDING = re.compile(r"待完善|待提供|待更新|待填|待补充")
OK_SHORT = re.compile(
    r"^(概述|报废原因|正线|南延伸|北延伸|东延伸|西延伸|西西延伸|"
    r"国际标准|国家标准|评估对象|评估的内容|评估涵盖时间|各系统评分情况|"
    r"隐患排查手册)$"
)
# 4.4.1 变电所系统MTBF（R）这类小节标题行，不能当「缺句号」标黄
_NUM_SECTION_TITLE = re.compile(r"^\d+(?:[\.．]\d+)+\s*\S")
CLAUSE_SPLIT = re.compile(r"[^。！？；]+[。！？；]?")

# 成文时由各章 write 写入，供 add_para / needs_mark 做年份校验
_ASSESSMENT_YEAR: int | None = None


def set_assessment_year(year: int | None) -> None:
    global _ASSESSMENT_YEAR
    _ASSESSMENT_YEAR = int(year) if year else None


def get_assessment_year() -> int | None:
    return _ASSESSMENT_YEAR


def is_section_title_line(text: str) -> bool:
    """带编号的小节标题行（如 4.4.1 …），不是正文句子。"""
    t = (text or "").strip()
    if not t or len(t) > 64:
        return False
    if t[-1] in "。！？；":
        return False
    return bool(_NUM_SECTION_TITLE.match(t))


def _is_caption_text(text: str) -> bool:
    t = (text or "").strip()
    if len(t) < 2 or len(t) >= 80 or t[0] not in {"图", "表"}:
        return False
    return t[1].isdigit() or t[1] in " -"


def _unclosed(text: str) -> bool:
    return text.count("（") > text.count("）") or text.count("(") > text.count(")")


def is_incomplete(text: str) -> bool:
    """话说一半、待提供、缺关键数字。缺句末句号或双分号不算。

    「故障起」后面没有数字、段首「号线」「026年」、括号没闭合、收在「的和与及」上，都当不通顺。
    引导语（具体如下）和图题表题不当残句。
    """
    t = (text or "").strip()
    if not t or LEAD_IN.match(t) or _is_caption_text(t):
        return False
    packed = t.replace(" ", "")
    if FAULT_GAP.search(packed):
        return True
    if PENDING.search(t):
        return True
    if re.match(r"^026年", t) or re.match(r"^号线", t):
        return True
    if _unclosed(t):
        return True
    last = t[-1]
    if last in "。！？":
        return False
    if last in "：:；":
        return False
    if OK_SHORT.match(t):
        return False
    if CUT_END.search(t):
        return True
    return False


def split_clauses(text: str) -> list[str]:
    """按句号、问号、分号切开，标点留在该句末。整段只黄出错的那一句。"""
    t = str(text or "")
    if not t:
        return []
    return [p for p in CLAUSE_SPLIT.findall(t) if p]


def is_punct_issue(text: str) -> bool:
    """缺句末标点、双分号等。收在「退出运营」「运行情况」上但没有句号，也算。

    「——修订了…」清单条目在材料里偶有漏写句末分号，按材料原文照抄，不因此整句黄。
    """
    t = (text or "").strip()
    if (
        not t
        or is_incomplete(t)
        or LEAD_IN.match(t)
        or _is_caption_text(t)
        or OK_SHORT.match(t)
        or is_section_title_line(t)
    ):
        return False
    if "；；" in t or "；。" in t or "。；" in t:
        return True
    # 破折号修订清单：不因缺句末标点标黄
    if t.startswith("——") or t.startswith("—") or t.startswith("--"):
        return False
    last = t[-1]
    if last in "。！？：:；":
        return False
    return True


def needs_mark(text: str) -> bool:
    """这一句要不要黄：残句、标点、数字合计、年份逻辑。数字/年份按整段判断。"""
    from chapters.common.number_logic import has_number_logic_issue, has_year_logic_issue

    if is_section_title_line(text):
        return has_year_logic_issue(text, get_assessment_year())
    return (
        is_incomplete(text)
        or is_punct_issue(text)
        or has_number_logic_issue(text)
        or has_year_logic_issue(text, get_assessment_year())
    )


def any_incomplete(texts: list[str] | None) -> bool:
    """小节里有话说一半的句子。"""
    return any(is_incomplete(t) for t in texts or [] if str(t).strip())


def any_needs_mark(texts: list[str] | None) -> bool:
    """小节里有不通顺或标点问题，对应标题也要黄。"""
    return any(needs_mark(t) for t in texts or [] if str(t).strip())


def punct_note(text: str) -> str:
    """网页「标点错误」后面写清错在哪、为什么，不整段摘抄。"""
    t = (text or "").replace("\n", "").strip()
    if "；；" in t:
        return "句中出现连续分号「；；」"
    if "；。" in t:
        return "分号后面又紧跟句号「；。」"
    if "。；" in t:
        return "句号后面又紧跟分号「。；」"
    clauses = split_clauses(t) or [t]
    bad = next((c.strip() for c in clauses if is_punct_issue(c.strip())), clauses[-1].strip())
    clip = bad if len(bad) <= 24 else "…" + bad[-24:]
    return f"末句「{clip}」没有句末句号、问号或分号"


def language_note(text: str) -> str:
    """网页「语言不通顺」后面写清错在哪、为什么。"""
    t = (text or "").replace("\n", "").strip()
    packed = t.replace(" ", "")
    if FAULT_GAP.search(packed):
        return "「故障起」后面没有写出起数"
    if PENDING.search(t):
        return "句中有待提供、待完善等，材料未写完"
    if _unclosed(t):
        return "括号没有写完"
    if re.match(r"^026年", t):
        return "段首年份缺字（如「026年」）"
    if re.match(r"^号线", t):
        return "段首缺线路名，只剩「号线」"
    if CUT_END.search(t):
        return "句子收在顿号或「的、和、与」上，话说一半"
    clip = t if len(t) <= 24 else t[:24] + "…"
    return f"「{clip}」话说一半，未写完整"


def new_report_document() -> Document:
    """空白年报：宋体 12 磅、A4、2025 页边距。各章 write 都从这里开文档。"""
    doc = Document()
    _set_font(doc)
    for section in doc.sections:
        # 2025 供电年报：A4，上下 1 英寸，左右 1.25 英寸。
        section.page_width = Emu(7560310)
        section.page_height = Emu(10692130)
        section.top_margin = Emu(914400)
        section.bottom_margin = Emu(914400)
        section.left_margin = Emu(1143000)
        section.right_margin = Emu(1143000)
    return doc


def add_cover(doc: Document, *, year: int, cover_line: str) -> None:
    """封面两行：总报告名 + 本章/分册名。cover_line 由各章提供，不在这里写专业判断。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    _run(p, f"上海轨道交通运营设施设备{year}年度评估报告", bold=True, size=18, name=BODY_FONT)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.first_line_indent = Cm(0)
    _run(p2, cover_line, bold=True, size=16, name=BODY_FONT)


def add_heading(doc: Document, title: str, level: int, source: str = "", *, yellow: bool = False) -> None:
    """小节标题。有出处只拼在标题后括号里；yellow 黄整条标题（缺材料或该节有不通顺句）。"""
    text = f"{title}（{source}）" if source else title
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(min(max(level, 1), 9) - 1))
    p_pr.append(outline)
    _run(p, text, bold=False, size=BODY_PT, name=BODY_FONT, color=RGBColor(0, 0, 0), highlight=yellow)
    _apply_heading_paragraph(p)


def add_caption(doc: Document, text: str) -> None:
    """图题/表题居中、无首行缩进。不是正文，不标黄、不加出处括号。"""
    if not text:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(0)
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    _run(p, text, bold=False, size=BODY_PT, name=BODY_FONT)


def add_table(
    doc: Document,
    rows: list[list[str]],
    *,
    widths: list[int] | None = None,
    header_bold: bool | None = None,
    col_align: list[str] | None = None,
    jc: str | None = None,
    cell_mar: int | None = None,
    vmerge: list[tuple[int, int, int]] | None = None,
    font_pt: float | None = None,
    domain_id: str | None = None,
) -> None:
    """插表。供电套供电 2025 表头规格；接触网不用供电规格。"""
    if not rows:
        return
    from chapters.common.docx_style import BODY_PT as DEFAULT_PT
    from chapters.common.docx_style import _table

    kwargs = {
        "widths": widths,
        "header_bold": header_bold,
        "col_align": col_align,
        "jc": jc,
        "cell_mar": cell_mar,
        "vmerge": vmerge,
        "font_pt": font_pt,
    }
    if domain_id == "overhead":
        from chapters.common.table_layout import pad_widths

        spec = {k: v for k, v in kwargs.items() if v is not None}
        ncols = len(rows[0])
        if ncols:
            spec["widths"] = pad_widths(spec.get("widths"), ncols)
    else:
        from chapters.power.format import merge_table_kwargs

        spec = merge_table_kwargs(rows, kwargs)
    merges = spec.get("vmerge")
    if merges is None:
        merges = infer_vmerges(rows)
    _table(
        doc,
        [str(c) for c in rows[0]],
        [[str(c) for c in r] for r in rows[1:]],
        widths=spec.get("widths"),
        header_bold=bool(spec.get("header_bold")),
        col_align=spec.get("col_align"),
        jc=spec.get("jc", "center"),
        cell_mar=spec.get("cell_mar"),
        vmerge=merges,
        font_pt=spec.get("font_pt") if spec.get("font_pt") is not None else DEFAULT_PT,
    )


def _formula_run(p, text: str, *, vert: str | None = None) -> None:
    _run(p, text, size=BODY_PT, name=BODY_FONT, vert=vert)


def add_formula_3_1(doc: Document) -> None:
    """公式居中、编号（3-1）右对齐。用上下标写出，避免 OMML 在 WPS 里空白。"""
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(0)
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.tab_stops.add_tab_stop(Cm(14.65), WD_TAB_ALIGNMENT.RIGHT)
    _formula_run(p, "s = Σ")
    _formula_run(p, "i", vert="subscript")
    _formula_run(p, "N", vert="superscript")
    _formula_run(p, " w")
    _formula_run(p, "i", vert="subscript")
    _formula_run(p, " v")
    _formula_run(p, "i", vert="subscript")
    _formula_run(p, "\t（3-1）")

    add_para(doc, "其中：", indent=False)
    defs = (
        (("s", None), " —— 子系统线路状态评分；"),
        (("w", "i"), " —— 子系统中 i 类设备的权重；"),
        (("v", "i"), " —— 子系统中 i 类设备的评分；"),
    )
    for (base, sub), text in defs:
        line = doc.add_paragraph()
        _apply_body_paragraph(line)
        line.paragraph_format.first_line_indent = Cm(0)
        _formula_run(line, base)
        if sub:
            _formula_run(line, sub, vert="subscript")
        _formula_run(line, text)


def add_blank(doc: Document) -> None:
    """表后空一段，对应 2025 年报表 3-3 / 3-10 之间的间距，不是缺材料占位。"""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5


def add_para(
    doc: Document,
    text: str,
    *,
    yellow: bool = False,
    indent: bool = True,
    align: str | None = None,
    source_yellow: bool = False,
) -> None:
    """正文。数字/年份逻辑对整段黄；标点/残句只黄出错分句。小节标题行不因缺句号黄。"""
    if not text:
        return
    chunks = [p.strip() for p in str(text).split("\n") if p.strip()]
    if len(chunks) > 1:
        for chunk in chunks:
            add_para(
                doc,
                chunk,
                yellow=yellow,
                indent=indent,
                align=align,
                source_yellow=source_yellow,
            )
        return
    from chapters.common.number_logic import has_number_logic_issue, has_year_logic_issue
    from chapters.common.source_yellow import apply_source_yellow

    text, from_src = apply_source_yellow(text, source_yellow=source_yellow)
    if from_src:
        yellow = True

    p = doc.add_paragraph()
    _apply_body_paragraph(p)
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if yellow:
        _run(p, text, size=BODY_PT, name=BODY_FONT, highlight=True)
        return
    if is_section_title_line(text):
        # 4.4.1 这类当正文写出时也不要黄；年份问题仍黄
        mark = has_year_logic_issue(text, get_assessment_year())
        _run(p, text, size=BODY_PT, name=BODY_FONT, highlight=mark)
        return
    # 故障总分项等跨分号，必须整段判断，不能切开后丢「其中」
    if has_number_logic_issue(text) or has_year_logic_issue(text, get_assessment_year()):
        _run(p, text, size=BODY_PT, name=BODY_FONT, highlight=True)
        return
    clauses = split_clauses(text) or [text]
    for clause in clauses:
        _run(p, clause, size=BODY_PT, name=BODY_FONT, highlight=needs_mark(clause))


def source_name(path: str | Path | None, fallback: str = "") -> str:
    """标题括号里的出处：只用文件名，不写路径、不写〔出处〕脚注。"""
    if path:
        return Path(path).name
    return fallback
