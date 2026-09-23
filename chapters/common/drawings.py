# -*- coding: utf-8 -*-
"""把源 Word 里的图片、图表原样拷进报告，不另画、不编造。

关系 id 要重映射，否则两份文档会抢同一 rId。图幅限制在 2025 正文宽度内。
"""
from __future__ import annotations

import posixpath
import re
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn
from docx.shared import Cm

from chapters.power.format import FIGURE_MAX_CX, FIGURE_MIN_CX

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RID_ATTRS = {
    f"{{{R_NS}}}id",
    f"{{{R_NS}}}embed",
    f"{{{R_NS}}}link",
}


def drawing_flow_item(doc: Any | None, block: Any | None) -> dict[str, Any]:
    """抽取时记下源路径与 body 下标，成文按此回拷，不重绘。"""
    idx = -1
    if block is not None and getattr(block, "source_index", None) is not None:
        try:
            idx = int(block.source_index)
        except (TypeError, ValueError):
            idx = -1
    path = str(getattr(doc, "source_path", "") or "")
    item: dict[str, Any] = {"kind": "drawing"}
    if idx >= 0 and path.lower().endswith((".docx", ".doc")):
        item["source_path"] = path
        item["source_index"] = idx
    return item


def paragraph_has_drawing(paragraph_elm) -> bool:
    """该段是否带 drawing/pict/OLE。空段不要当图复制。"""
    if paragraph_elm is None:
        return False
    return bool(
        paragraph_elm.findall(".//" + qn("w:drawing"))
        or paragraph_elm.findall(".//" + qn("w:pict"))
        or paragraph_elm.findall(".//" + qn("w:object"))
        or paragraph_elm.findall(".//{urn:schemas-microsoft-com:vml}imagedata")
    )


def _drawing_nodes(src_elm):
    nodes = []
    seen: set[int] = set()
    for tag in (qn("w:drawing"), qn("w:pict"), qn("w:object")):
        for node in src_elm.findall(".//" + tag):
            key = id(node)
            if key in seen:
                continue
            seen.add(key)
            nodes.append(node)
    return nodes


def _rids_in(node) -> set[str]:
    needed: set[str] = set()
    for el in node.iter():
        for attr, val in el.attrib.items():
            local = attr.split("}")[-1]
            if local in {"id", "embed", "link"} and str(val).startswith("rId"):
                needed.add(val)
    return needed


def _ensure_ole_shapetype(src_doc, cloned) -> None:
    """WPS/Word 的 OLE 对象依赖文档级 shapetype；只拷 object 节点会变成空白框。"""
    if cloned.tag != qn("w:object"):
        return
    if any(el.tag.endswith("shapetype") for el in cloned.iter()):
        return
    for el in src_doc.element.body.iter():
        if el.tag.endswith("shapetype") and el.get("id") == "_x0000_t75":
            cloned.insert(0, deepcopy(el))
            return


def _unique_partname(package, original: PackURI) -> PackURI:
    folder = original.baseURI
    stem, ext = posixpath.splitext(original.filename)
    base = re.sub(r"\d+$", "", stem) or stem
    template = f"{folder}/{base}%d{ext}"
    return package.next_partname(template)


def _clone_part(src_part: Part, dest_package, cache: dict[int, Part]) -> Part:
    """把图片/OLE 二进制拷进目标包。同一源 part 只建一次，避免一张图嵌入两次。"""
    key = id(src_part)
    if key in cache:
        return cache[key]
    partname = _unique_partname(dest_package, src_part.partname)
    new_part = Part(partname, src_part.content_type, src_part.blob, dest_package)
    cache[key] = new_part
    for rel in list(src_part.rels.values()):
        if rel.is_external:
            new_part.load_rel(rel.reltype, rel.target_ref, rel.rId, True)
            continue
        child = _clone_part(rel.target_part, dest_package, cache)
        new_part.load_rel(rel.reltype, child, rel.rId, False)
    return new_part


def anchor_to_inline(node):
    """浮动图（wp:anchor + wrap）改成嵌入图，避免后文绕到图旁边。"""
    tag = getattr(node, "tag", "") or ""
    if not str(tag).endswith("}anchor"):
        return node
    ns = str(tag).split("}")[0][1:]
    inline = node.makeelement(f"{{{ns}}}inline")
    for key, val in (("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0")):
        inline.set(key, val)
    keep = {"extent", "effectExtent", "docPr", "cNvGraphicFramePr", "graphic"}
    for child in list(node):
        local = str(child.tag).split("}")[-1]
        if local in keep:
            inline.append(child)
    parent = node.getparent()
    if parent is not None:
        parent.replace(node, inline)
    return inline


def inline_floating_drawings(root):
    """w:drawing 里包着的 wp:anchor 也要改成嵌入。"""
    if root is None:
        return root
    if str(getattr(root, "tag", "")).endswith("}anchor"):
        root = anchor_to_inline(root)
    for el in [x for x in root.iter() if str(getattr(x, "tag", "")).endswith("}anchor")]:
        anchor_to_inline(el)
    return root


def scale_drawing_node(node, max_cx: int = FIGURE_MAX_CX, min_cx: int = FIGURE_MIN_CX) -> None:
    """把图幅限制在 2025 年报正文宽度内，过小的装饰符不放大。"""
    if node is None:
        return
    cx = cy = None
    targets = []
    for el in node.iter():
        if el.get("cx") and el.get("cy") and (el.tag.endswith("ext") or el.tag.endswith("extent")):
            targets.append(el)
            if cx is None:
                try:
                    cx = int(el.get("cx"))
                    cy = int(el.get("cy"))
                except (TypeError, ValueError):
                    pass
    if not cx or not cy or cx <= min_cx or cx <= max_cx:
        return
    ratio = max_cx / cx
    new_cx = str(int(cx * ratio))
    new_cy = str(int(cy * ratio))
    for el in targets:
        el.set("cx", new_cx)
        el.set("cy", new_cy)


def _remap_rids(elm, rid_map: dict[str, str]) -> None:
    """把源文档 rId 换成目标文档新 id，否则图会指到空关系。"""
    for node in elm.iter():
        for attr in list(node.attrib):
            if attr not in RID_ATTRS:
                continue
            old = node.get(attr)
            if old in rid_map:
                node.set(attr, rid_map[old])


def _docx_for_copy(src_path: str | Path) -> Path | None:
    """.doc 不能直接给 python-docx；优先已转换的 sidecar，没有再转一次。"""
    src_path = Path(src_path)
    if src_path.suffix.lower() == ".docx" and src_path.is_file():
        return src_path
    if src_path.suffix.lower() != ".doc":
        return None
    from parsers.office_convert import converted_sidecar, convert_legacy

    hit = converted_sidecar(src_path, None, "docx")
    if hit is not None:
        return hit
    if not src_path.is_file():
        return None
    try:
        return convert_legacy(src_path, src_path.parent / "_converted")
    except Exception:
        return None


def copy_drawings_from_body_index(src_path: str | Path, source_index: int, dest_doc: Document) -> bool:
    """按源文档 body 子节点下标拷图。抽取时记下 source_index，成文时原样贴回，不重绘。"""
    if source_index < 0:
        return False
    resolved = _docx_for_copy(src_path)
    if resolved is None:
        return False
    src = Document(str(resolved))
    children = list(src.element.body.iterchildren())
    if source_index >= len(children):
        return False
    src_elm = children[source_index]
    if not paragraph_has_drawing(src_elm):
        return False
    # 证书/登记表常嵌在单列表里；整表回拷才能保住原尺寸。
    if src_elm.tag == qn("w:tbl"):
        from chapters.common.table_copy import copy_table_from_body_index

        return copy_table_from_body_index(resolved, source_index, dest_doc)
    nodes = _drawing_nodes(src_elm)
    if not nodes:
        return False
    cache: dict[int, Part] = {}
    copied = False
    for node in nodes:
        dest_p = dest_doc.add_paragraph()
        dest_p.paragraph_format.first_line_indent = Cm(0)
        dest_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rid_map: dict[str, str] = {}
        needed = _rids_in(node)
        for rid in needed:
            if rid not in src.part.rels:
                continue
            rel = src.part.rels[rid]
            if rel.is_external:
                new_rid = dest_doc.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                new_part = _clone_part(rel.target_part, dest_doc.part.package, cache)
                new_rid = dest_doc.part.relate_to(new_part, rel.reltype)
            rid_map[rid] = new_rid
        cloned = deepcopy(node)
        _ensure_ole_shapetype(src, cloned)
        _remap_rids(cloned, rid_map)
        scale_drawing_node(cloned)
        cloned = inline_floating_drawings(cloned)
        run = dest_p.add_run()
        run._r.append(cloned)
        copied = True
    return copied


M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def paragraph_has_omml(paragraph_elm) -> bool:
    if paragraph_elm is None:
        return False
    return bool(
        paragraph_elm.findall(f".//{{{M_NS}}}oMath")
        or paragraph_elm.findall(f".//{{{M_NS}}}oMathPara")
    )


def copy_formula_paragraph_from_body_index(src_path: str | Path, source_index: int, dest_doc: Document) -> bool:
    """按 source_index 把源段整段拷进报告（含 OMML 公式与段内嵌入对象），保持公式形态。"""
    if source_index < 0:
        return False
    resolved = _docx_for_copy(src_path)
    if resolved is None:
        return False
    src = Document(str(resolved))
    children = list(src.element.body.iterchildren())
    if source_index >= len(children):
        return False
    src_elm = children[source_index]
    if src_elm.tag != qn("w:p"):
        return False
    if not (paragraph_has_omml(src_elm) or paragraph_has_drawing(src_elm)):
        return False

    cloned = deepcopy(src_elm)
    cache: dict[int, Part] = {}
    rid_map: dict[str, str] = {}
    needed = _rids_in(cloned)
    for rid in needed:
        if rid not in src.part.rels:
            continue
        rel = src.part.rels[rid]
        if rel.is_external:
            rid_map[rid] = dest_doc.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            new_part = _clone_part(rel.target_part, dest_doc.part.package, cache)
            rid_map[rid] = dest_doc.part.relate_to(new_part, rel.reltype)
    for node in _drawing_nodes(cloned):
        _ensure_ole_shapetype(src, node)
        # 公式旁小嵌入对象不按大图缩放
        if not paragraph_has_omml(src_elm):
            scale_drawing_node(node)
    _remap_rids(cloned, rid_map)
    # 用 add_paragraph 占位再替换，保证插在当前成文位置，而不是 body 末尾。
    placeholder = dest_doc.add_paragraph()
    parent = placeholder._element.getparent()
    idx = list(parent).index(placeholder._element)
    parent.remove(placeholder._element)
    parent.insert(idx, cloned)
    return True
