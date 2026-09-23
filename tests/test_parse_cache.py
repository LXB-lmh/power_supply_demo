# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document

from parsers.dispatch import parse_files
from parsers.parse_cache import (
    _resolved_source_path,
    load_classify_latest,
    load_document,
    parse_file_cached,
    persist_classify_materials,
    public_cache_info,
    save_classify_result,
)


def test_resolved_source_path_keeps_converted_docx(tmp_path: Path):
    original = tmp_path / "去年.doc"
    original.write_bytes(b"doc")
    converted = tmp_path / "_converted"
    converted.mkdir()
    docx = converted / "去年.docx"
    Document().save(str(docx))
    got = Path(_resolved_source_path(original, str(original), None))
    assert got.resolve() == docx.resolve()


def test_parse_file_cached_roundtrip(tmp_path: Path):
    src = tmp_path / "a.docx"
    doc = Document()
    doc.add_paragraph("缓存测试段落")
    doc.save(str(src))

    d1, dig1, hit1 = parse_file_cached(src)
    assert not hit1
    assert "缓存测试" in d1.to_text()

    d2, dig2, hit2 = parse_file_cached(src)
    assert hit2
    assert dig1 == dig2
    assert load_document(dig1) is not None

    docs = parse_files([src])
    assert len(docs) == 1


def test_classify_cache_latest(tmp_path: Path):
    src = tmp_path / "a.docx"
    Document().save(str(src))
    persist_classify_materials("overhead", [src], ["a.docx"])
    save_classify_result(
        "overhead",
        result={"total": 1, "assigned": 1, "by_chapter": {"ch3": [{"relpath": "a.docx", "name": "a.docx"}]}},
        files_meta=[{"rel": "a.docx", "content_sha256": "abc", "name": "a.docx"}],
        assessment_year=2026,
    )
    row = load_classify_latest("overhead")
    assert row and row.get("result")
    info = public_cache_info("overhead")
    assert info.get("available")
    assert info.get("stored_count") == 1


def test_parse_files_cached_parses_duplicate_once(tmp_path: Path, monkeypatch):
    import shutil

    from parsers import parse_cache as pc

    src = tmp_path / "a.docx"
    dup = tmp_path / "copy.docx"
    doc = Document()
    doc.add_paragraph("同一份材料复制两遍")
    doc.save(str(src))
    shutil.copy2(src, dup)

    calls = {"n": 0}
    real = pc.parse_file_cached

    def wrapped(path, work_dir=None):
        calls["n"] += 1
        return real(path, work_dir=work_dir)

    monkeypatch.setattr(pc, "parse_file_cached", wrapped)
    docs, stats = pc.parse_files_cached([src, dup])
    assert stats.get("unique_files") == 1
    assert calls["n"] == 1
    assert len(docs) == 2
    assert docs[0] is not docs[1]
    assert docs[0].source_name == "a.docx"
    assert docs[1].source_name == "copy.docx"
