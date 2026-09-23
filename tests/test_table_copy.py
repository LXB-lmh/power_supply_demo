# -*- coding: utf-8 -*-
"""Word 表格整表回拷（颜色/列宽）。"""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from chapters.common.table_copy import (
    WEIGHT_RULE_WIDTHS_7,
    copy_table_from_body_index,
    resolve_docx_for_copy,
    table_flow_item,
    write_table_item,
)
from chapters.common.word import new_report_document
from parsers.docx_parser import parse_docx


def _shade_cell(cell, fill: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def test_copy_table_preserves_cell_shading(tmp_path: Path):
    src = tmp_path / "src.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "状态参数"
    table.cell(0, 1).text = "权重"
    _shade_cell(table.cell(0, 0), "4472C4")
    _shade_cell(table.cell(0, 1), "4472C4")
    table.cell(1, 0).text = "接触线磨耗"
    table.cell(1, 1).text = "0.2"
    doc.save(str(src))

    parsed = parse_docx(src)
    block = next(b for b in parsed.blocks if b.type == "table")
    item = table_flow_item(parsed, block)
    assert item.get("source_index") == 0
    assert item.get("source_path")

    out = new_report_document()
    assert write_table_item(out, item)
    out_path = tmp_path / "out.docx"
    out.save(str(out_path))

    out_tbl = Document(str(out_path)).tables[0]
    shd = out_tbl.cell(0, 0)._tc.tcPr.find(qn("w:shd"))
    assert shd is not None
    assert shd.get(qn("w:fill")) == "4472C4"


def test_copy_table_from_body_index_direct(tmp_path: Path):
    src = tmp_path / "t.docx"
    doc = Document()
    doc.add_paragraph("前段")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "唯一表"
    doc.save(str(src))
    dest = new_report_document()
    assert copy_table_from_body_index(src, 1, dest)
    assert len(dest.tables) == 1
    assert dest.tables[0].cell(0, 0).text == "唯一表"


def test_copy_table_resolves_doc_sidecar(tmp_path: Path):
    converted_dir = tmp_path / "_converted"
    converted_dir.mkdir()
    src = converted_dir / "prior.docx"
    doc = Document()
    doc.add_paragraph("前段")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "状态参数"
    table.cell(0, 1).text = "权重"
    doc.save(str(src))
    dummy = tmp_path / "prior.doc"
    dummy.write_bytes(b"not-a-real-doc")
    assert resolve_docx_for_copy(dummy).resolve() == src.resolve()
    dest = new_report_document()
    assert copy_table_from_body_index(dummy, 1, dest)
    assert dest.tables[0].cell(0, 0).text.strip() == "状态参数"


def test_weight_rule_table_rebuild_uses_last_year_col_widths():
    dest = new_report_document()
    rows = [
        ["状态参数", "权重", "评分计算规则", "", "状态参数", "权重", "评分计算规则"],
        ["接触线磨耗", "5", "优：5分", "", "接触线导高情况", "5", "优：5分"],
    ]
    assert write_table_item(dest, {"kind": "table", "rows": rows, "force_rebuild": True})
    grid = dest.tables[0]._tbl.find(qn("w:tblGrid"))
    cols = [int(c.get(qn("w:w"))) for c in grid.findall(qn("w:gridCol"))]
    assert cols == WEIGHT_RULE_WIDTHS_7
    assert cols[2] > cols[0] * 3
    assert cols[3] < cols[0]


def test_weight_rule_table_rebuild_uses_prior_colors(tmp_path: Path):
    dest = new_report_document()
    rows = [
        ["状态参数", "权重", "评分计算规则", "", "状态参数", "权重", "评分计算规则"],
        ["接触线磨耗", "5", "优：5分", "", "接触线导高情况", "5", "优：5分"],
        ["接触线线面状态", "5", "优：5分", "", "接地系统状态", "5", "优：5分"],
    ]
    assert write_table_item(dest, {"kind": "table", "rows": rows, "force_rebuild": True})
    tbl = dest.tables[0]
    shd = tbl.cell(0, 0)._tc.tcPr.find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == "2F5597"
    shd1 = tbl.cell(1, 0)._tc.tcPr.find(qn("w:shd"))
    assert shd1 is not None and shd1.get(qn("w:fill")) == "D2DEEF"


def test_status_dual_header_merges_abcd_pairs():
    dest = new_report_document()
    rows = [
        ["线路", "锚段数量", "A", "A", "B", "B", "C", "C", "D", "D"],
        ["线路", "锚段数量", "锚段数", "占比", "锚段数", "占比", "锚段数", "占比", "锚段数", "占比"],
        ["1号线", "179", "111", "62.01%", "18", "10.06%", "50", "27.93%", "0", "0.00%"],
    ]
    from chapters.common.table_copy import add_table_full_rows

    add_table_full_rows(dest, rows)
    tbl = dest.tables[0]
    assert tbl.cell(0, 2)._tc is tbl.cell(0, 3)._tc
    assert (tbl.cell(0, 2).text or "").strip() == "A"
    assert tbl.cell(0, 4)._tc is tbl.cell(0, 5)._tc
    assert (tbl.cell(0, 4).text or "").strip() == "B"
    assert tbl.cell(0, 6)._tc is tbl.cell(0, 7)._tc
    assert tbl.cell(0, 8)._tc is tbl.cell(0, 9)._tc
    assert tbl.cell(1, 2)._tc is not tbl.cell(1, 3)._tc
    assert (tbl.cell(1, 2).text or "").strip() == "锚段数"
    assert (tbl.cell(1, 3).text or "").strip() == "占比"
    assert tbl.cell(0, 0)._tc is tbl.cell(1, 0)._tc
    assert tbl.cell(0, 1)._tc is tbl.cell(1, 1)._tc


def test_cycle_maintain_table_uses_human_widths():
    from chapters.common.table_copy import collapse_cycle_maintain_tables, write_table_item
    from chapters.overhead.style import OH_CYCLE_4COL_WIDTHS

    dest = new_report_document()
    rows = [
        ["大类", "工作项目", "维护周期", "维护内容"],
        ["巡视", "柔性接触网", "三个月一次", "步行巡视（段场）"],
        ["", "", "六个月一次", "步行巡视（隧道段）"],
    ]
    assert write_table_item(dest, {"kind": "table", "rows": rows})
    tbl = dest.tables[0]
    grid = [g.get(qn("w:w")) for g in tbl._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
    assert grid == [str(w) for w in OH_CYCLE_4COL_WIDTHS]
    mar = tbl._tbl.tblPr.find(qn("w:tblCellMar")) if tbl._tbl.tblPr is not None else None
    assert mar is None
    flow = [
        {"kind": "table", "rows": rows[:2]},
        {"kind": "para", "text": "接触轨包括：16正线。"},
        {"kind": "table", "rows": rows},
    ]
    kept = collapse_cycle_maintain_tables(flow)
    assert sum(1 for x in kept if x.get("kind") == "table") == 1
    assert len(kept[1]["rows"]) == 3
