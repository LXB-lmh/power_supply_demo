# -*- coding: utf-8 -*-
"""小节审阅口径：缺材料、逻辑不通、语言不通顺、标点错误分条；空措施不是缺文件。"""
from pathlib import Path

from chapters.power.extract import extract_compliance
from chapters.power.review import review_chapter
from chapters.ch4.power_extract import build_ch4_pack
from parsers.dispatch import parse_file, parse_files

PACK = Path(r"F:\材料\2026供电评估_新\评估材料\99-26年评估报告编制材料")
COMP = PACK / "4&5合规性材料(2026).docx"
FAULT = PACK / "评估材料（故障情况，黄色待更新）.docx"
ZONG = PACK / "设备功能有效性" / "设备评估结果总表（4月）_1.xlsx"


def test_ch5_review_marks_empty_sections():
    if not COMP.exists():
        return
    pack = extract_compliance([parse_file(COMP)])
    pack["chapter_id"] = "ch5"
    rows = review_chapter(pack, "ch5")
    issues = {(x["section"], x["issue"]) for x in rows}
    assert ("5.1", "缺失材料") in issues
    assert ("5.2", "缺失材料") in issues
    assert ("5.7", "缺失材料") in issues  # 评估小结一律固定缺
    assert ("5.3", "缺失材料") not in issues  # 合规性材料里这两节已有正文
    assert ("5.6", "缺失材料") not in issues


def test_ch6_review_logic_and_summary():
    if not COMP.exists():
        return
    pack = extract_compliance([parse_file(COMP)])
    rows = review_chapter(pack, "ch6")
    issues = {(x["section"], x["issue"]) for x in rows}
    assert ("6.3", "缺失材料") in issues  # 评估小结一律固定缺
    assert ("6.2", "缺失材料") not in issues
    assert ("6.1", "缺失材料") not in issues
    assert ("6.1", "逻辑不通") in issues  # 修程修制正文与表对不上，不是缺文件


def test_ch4_review_language_and_mtbf():
    if not FAULT.exists() or not ZONG.exists():
        return
    pack = build_ch4_pack(parse_files([ZONG, FAULT]), year=2026)
    rows = review_chapter(pack, "ch4")
    issues = {(x["section"], x["issue"]) for x in rows}
    assert ("4.4", "缺失材料") in issues  # 无生产指标材料：只记 4.4，不写死 4.4.1～4.4.5
    assert ("4.5", "缺失材料") in issues  # 评估小结一律固定缺
    assert ("4.1.2", "缺失材料") not in issues  # 总表已够 4.1.2
    assert any(s.startswith("4.3") and i == "语言不通顺" for s, i in issues)


def test_ch3_review_only_when_no_tables():
    pack = {"ch3": {}}
    rows = review_chapter(pack, "ch3")
    issues = {x["section"] for x in rows if x["issue"] == "缺失材料"}
    assert "3.1" in issues
    assert "3.2" in issues
    assert "3.3.1" in issues
    assert "3.3.2" in issues
    assert "3.4" in issues
    assert "3.6" in issues


def test_ch3_punct_yellows_bad_clause_and_titles(tmp_path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from chapters.power.write import write_chapter_docx

    pack = {
        "year": 2026,
        "ch3": {
            "controls": {
                "source": "0管控措施（设备树）.xlsx",
                "prose": [
                    {
                        "line": "2号线",
                        "blocks": ["正线：\n变压器设备：暂未纳入大修更新改造规划；暂无备件"],
                    }
                ],
            }
        },
    }
    out = tmp_path / "ch3.docx"
    write_chapter_docx(pack, out, "ch3")
    doc = Document(str(out))

    def yellow(p):
        return any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in p.runs)

    titles = {p.text.strip(): yellow(p) for p in doc.paragraphs if p.text.strip()}
    assert titles.get("轨道交通2号线") is True  # 该线有标点问题，线路标题也黄
    assert titles.get("正线：") is True  # 区段标题一并黄
    assert titles.get("轨道交通1号线") is True  # 该线无管控正文，标黄标题
    sent = next(p for p in doc.paragraphs if "暂无备件" in p.text)
    assert yellow(sent) is True
    # 只黄缺句号的那一分句，前面带分号的不黄
    assert sent.runs[-1].font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert "暂无备件" in sent.runs[-1].text
    assert sent.runs[0].font.highlight_color != WD_COLOR_INDEX.YELLOW


def test_ch3_review_lists_punct_and_keeps_missing_lines():
    pack = {
        "ch3": {
            "power_table": [["线路", "区段"], ["1号线", "正线"]],
            "compare_power": [["线路", "区段", "时间"], ["1号线", "正线", "2026"]],
            "migrations": ["1号线正线，应急电源设备，由A迁移到B"],
            "controls": {
                "prose": [
                    {
                        "line": "1号线",
                        "blocks": [
                            "正线：\n应急电源设备：暂未纳入大修更新改造规划；加强巡视，及时修复故障\n杂散电流设备：备件充裕；；车站短路器动作频繁。"
                        ],
                    }
                ]
            },
        }
    }
    rows = review_chapter(pack, "ch3")
    issues = {(x["section"], x["issue"]) for x in rows}
    assert ("3.4 1号线", "标点错误") in issues
    notes = [x.get("note") or "" for x in rows if x.get("issue") == "标点错误"]
    assert any("；；" in n or "没有句末" in n for n in notes)
    assert ("3.4 14号线", "缺失材料") in issues
    assert ("3.3.2", "缺失材料") not in issues  # 已有总表就不报 3.3.2 缺材料
    assert ("3.6", "缺失材料") in issues  # 评估小结一律固定缺


def test_ch3_review_empty_measures_not_missing_file():
    pack = {
        "ch3": {
            "power_table": [["线路"]],
            "compare_power": [["线路"]],
            "migrations": ["x"],
            "controls": {
                "prose": [{"line": "1号线", "blocks": ["正线：\n应急电源设备：备件充足。"]}],
                "appendix": [
                    ["线路", "区段", "大类", "中类", "评估结果", "大修", "差异化"],
                    ["14号线", "正线", "供电", "应急电源", "B", "", ""],
                    ["14号线", "南延伸", "供电", "变压器", "B", "/", "无"],
                ],
            },
        }
    }
    rows = review_chapter(pack, "ch3")
    issues = {(x["section"], x["issue"]) for x in rows}
    assert ("3.4 14号线", "措施未填") in issues  # 附录有行但措施空，不是缺文件
    assert ("3.4 14号线", "缺失材料") not in issues
    assert ("3.4 15号线", "缺失材料") in issues
