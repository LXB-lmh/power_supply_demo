# -*- coding: utf-8 -*-
"""逻辑性检测页：供电/触网、章节切片、跨章红标。"""
from docx import Document

from chapters.common.logic_page_check import _rule_paragraph_sum, run_logic_page_check


def _save(tmp_path, paragraphs):
    src = tmp_path / "sample.docx"
    doc = Document()
    for text, heading in paragraphs:
        if heading:
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    doc.save(str(src))
    return src


def _kinds(hit):
    return [f.get("kind") for f in hit["findings"]]


def test_rule_paragraph_sum_mismatch():
    notes = _rule_paragraph_sum("本类设备共有10种，其中甲类有4种，乙类有5种。")
    assert notes
    assert "10" in notes[0] and "9" in notes[0]


def test_rule_paragraph_sum_user_example():
    notes = _rule_paragraph_sum("我们评估了A、B设备共150台，其中A100台，B20台。")
    assert notes
    assert "150" in notes[0] and "120" in notes[0]


def test_rule_paragraph_sum_ok():
    notes = _rule_paragraph_sum("本类设备共有9种，其中甲类有4种，乙类有5种。")
    assert notes == []


def test_rule_paragraph_sum_ignores_repeated_total():
    notes = _rule_paragraph_sum("设备共120台，其中A100台，B20台，合计120台。")
    assert notes == []


def test_year_flags_cutoff_2025_and_2024(tmp_path):
    src = _save(
        tmp_path,
        [
            ("第4章 运营契合满意度评估", True),
            ("本次统计截止到2025年，另一处写截止到2024。", False),
            ("口径截至2026年，与评估年一致。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="power_supply")
    years = [f for f in hit["findings"] if f["kind"] == "year"]
    blob = " ".join(f["note"] for f in years)
    assert "截止到2025" in blob and "截止到2024" in blob
    assert "截至2026" not in blob and "截止到2026" not in blob
    assert all("第4章" in f["section"] for f in years)
    assert hit["marked_report"] == ""
    assert hit["marked_url"] == ""


def test_completion_date_not_flagged_as_year(tmp_path):
    """「截止2025年M月D日完成进度100%」是项目完工日期（评估周期 2025.5-2026.4 内），不报年份偏早。"""
    src = _save(
        tmp_path,
        [
            ("第4章 运营契合满意度评估", True),
            ("11号线吊弦专项整治共需完成510套，截止2025年11月3日完成进度100%。", False),
            ("截至2025年6月底数据，共有5个供电专业可靠性指标。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="power_supply")
    years = [f for f in hit["findings"] if f["kind"] == "year"]
    blob = " ".join(f["note"] for f in years)
    assert "完成进度" not in blob
    assert "截至2025" in blob  # 数据截止时点偏早仍报

def test_abbreviated_year_at_start_not_missing_digit():
    """段首「26年度」是合法两位简写年份，不是「026年」缺字。"""
    from chapters.common.word import is_incomplete

    assert not is_incomplete("26年度开展人员培训包含岗位能力评估及考核共计83人次。")
    assert is_incomplete("026年度开展人员培训共计83人次。")

def test_chapter_slice_ignores_other_chapters(tmp_path):
    src = _save(
        tmp_path,
        [
            ("第2章 概述", True),
            ("概述资料截止到2020年。", False),
            ("第3章 设备功能有效性评估", True),
            ("我们评估了A、B设备共150台，其中A100台，B20台。", False),
            ("第4章 运营契合满意度评估", True),
            ("统计截止到2024年。", False),
            ("第11章 退运报废倾向性评估", True),
            ("4号线隔离开关应报废。", False),
        ],
    )
    only3 = run_logic_page_check(
        src,
        assessment_year=2026,
        mode="chapter",
        chapter_id="ch3",
        domain_id="overhead",
    )
    assert any(f["kind"] == "logic" for f in only3["findings"])
    assert not any(f["kind"] == "year" for f in only3["findings"])
    assert not any(f["kind"] == "severe" for f in only3["findings"])
    assert all("第3章" in f["section"] for f in only3["findings"])

    missing = run_logic_page_check(
        src,
        assessment_year=2026,
        mode="chapter",
        chapter_id="ch9",
        domain_id="overhead",
    )
    assert missing["finding_count"] == 0
    assert "第9章" in missing["message"]


def test_same_chapter_later_scrap_is_language_not_severe(tmp_path):
    src = _save(
        tmp_path,
        [
            ("第3章 设备功能有效性评估", True),
            ("4号线汇流排评为A类。", False),
            ("随后确认4号线汇流排应报废。", False),
            ("附录A 库存", True),
            ("附录统计截止到2019年。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="chapter", chapter_id="ch3", domain_id="overhead")
    notes = " ".join(f["note"] for f in hit["findings"] if f["kind"] == "language")
    assert "汇流排" in notes
    assert not any(f["kind"] == "severe" for f in hit["findings"])
    assert not any("2019" in (f.get("note") or "") for f in hit["findings"])

    full = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    assert not any("2019" in (f.get("note") or "") for f in full["findings"])


def test_full_report_cross_chapter_severe_and_skips_early_chapters(tmp_path):
    src = _save(
        tmp_path,
        [
            ("第2章 概述", True),
            ("概述资料截止到2020年。", False),
            ("第3章 设备功能有效性评估", True),
            ("4号线隔离开关评为A类。", False),
            ("5号线接触线评为A类。", False),
            ("第11章 退运报废倾向性评估", True),
            ("4号线隔离开关应报废。", False),
            ("5号线隔离开关无需报废。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    severe = [f for f in hit["findings"] if f["kind"] == "severe"]
    assert len(severe) == 1
    assert "隔离开关" in severe[0]["note"]
    assert "4号线" in severe[0]["note"]
    assert "接触线" not in severe[0]["note"]
    assert not any("2020" in (f.get("note") or "") for f in hit["findings"])


def test_domain_devices_do_not_cross(tmp_path):
    src = _save(
        tmp_path,
        [
            ("第3章 设备功能有效性评估", True),
            ("2号线变压器评为A类。", False),
            ("2号线接触线评为A类。", False),
            ("第11章 退运报废倾向性评估", True),
            ("2号线变压器应报废。", False),
            ("2号线接触线应报废。", False),
        ],
    )
    power = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="power_supply")
    overhead = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    power_notes = " ".join(f["note"] for f in power["findings"] if f["kind"] == "severe")
    overhead_notes = " ".join(f["note"] for f in overhead["findings"] if f["kind"] == "severe")
    assert "变压器" in power_notes and "接触线" not in power_notes
    assert "接触线" in overhead_notes and "变压器" not in overhead_notes


def test_scrap_trend_sentence_is_not_severe(tmp_path):
    src = _save(
        tmp_path,
        [
            ("第3章 设备功能有效性评估", True),
            ("4号线隔离开关评为A类。", False),
            ("第11章 退运报废倾向性评估", True),
            ("本章对4号线隔离开关开展退运报废评估。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    assert not any(f["kind"] == "severe" for f in hit["findings"])


def test_typo_punct_language(tmp_path):
    src = _save(
        tmp_path,
        [
            ("第8章 风险隐患闭环度评估", True),
            ("导线松驰，需要安排复测。", False),
            ("接触线驰度超出允许值。", False),
            ("本次评估已经完成；；请复核。", False),
            ("本月故障起。", False),
        ],
    )
    overhead = run_logic_page_check(src, assessment_year=2026, mode="chapter", chapter_id="ch8", domain_id="overhead")
    kinds = set(_kinds(overhead))
    assert "typo" in kinds and "punct" in kinds and "language" in kinds
    typo_notes = " ".join(f["note"] for f in overhead["findings"] if f["kind"] == "typo")
    assert "松弛" in typo_notes and "弛度" in typo_notes

    power = run_logic_page_check(src, assessment_year=2026, mode="chapter", chapter_id="ch8", domain_id="power_supply")
    power_typo = " ".join(f["note"] for f in power["findings"] if f["kind"] == "typo")
    assert "松弛" in power_typo
    assert "弛度" not in power_typo


def test_single_chapter_file_without_heading(tmp_path):
    src = _save(tmp_path, [("导线松驰，需要安排复测。", False)])
    hit = run_logic_page_check(
        src,
        assessment_year=2026,
        mode="chapter",
        chapter_id="ch10",
        domain_id="overhead",
    )
    assert hit["finding_count"] >= 1
    assert all("第10章" in f["section"] for f in hit["findings"])


def test_toc_entries_are_not_findings(tmp_path):
    """Word 目录条目不查标点/语言/年份，也不当章标题。"""
    src = _save(
        tmp_path,
        [
            ("目录", False),
            ("第3章 设备功能有效性评估 102", False),
            ("3.1 评估范围 102", False),
            ("11.3 评估小结 128", False),
            ("第11章 退运报废倾向性评估 130", False),
            ("12 总结与建议 132", False),
            ("第3章 设备功能有效性评估", True),
            ("本次评估了A、B设备共150台，其中A100台，B20台。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    # 目录条目不产生任何错误（正文的错字、计算错误仍要查出）
    for f in hit["findings"]:
        assert " 102" not in f["excerpt"] and " 128" not in f["excerpt"] and " 132" not in f["excerpt"]
    assert any(f["kind"] == "logic" for f in hit["findings"])


def test_toc_chapter_entry_does_not_shift_current(tmp_path):
    """目录里的「第3章」条目不把后面的正文归属到错误章号。"""
    src = _save(
        tmp_path,
        [
            ("目录", False),
            ("第3章 设备功能有效性评估 102", False),
            ("第3章 设备功能有效性评估", True),
            ("导线松驰，需要安排复测。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    typo = [f for f in hit["findings"] if f["kind"] == "typo"]
    assert typo
    assert all("第3章" in f["section"] for f in typo)


def test_table_total_mismatch_is_logic_error(tmp_path):
    """表格合计行与分项求和不符，报计算错误。"""
    src = tmp_path / "table_no_total.docx"
    doc = Document()
    doc.add_heading("第8章 风险隐患闭环度评估", level=1)
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "线路"
    table.rows[0].cells[1].text = "故障起数"
    table.rows[0].cells[2].text = "备注"
    table.rows[1].cells[0].text = "2号线"
    table.rows[1].cells[1].text = "5"
    table.rows[2].cells[0].text = "3号线"
    table.rows[2].cells[1].text = "8"
    doc.save(str(src))
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="power_supply")
    assert not any(f["kind"] == "logic" and "合计" in f["note"] for f in hit["findings"])

    src2 = tmp_path / "table_total.docx"
    doc2 = Document()
    doc2.add_heading("第8章 风险隐患闭环度评估", level=1)
    t2 = doc2.add_table(rows=4, cols=3)
    t2.rows[0].cells[0].text = "线路"
    t2.rows[0].cells[1].text = "故障起数"
    t2.rows[0].cells[2].text = "备注"
    t2.rows[1].cells[0].text = "2号线"
    t2.rows[1].cells[1].text = "5"
    t2.rows[2].cells[0].text = "3号线"
    t2.rows[2].cells[1].text = "8"
    t2.rows[3].cells[0].text = "合计"
    t2.rows[3].cells[1].text = "10"
    doc2.save(str(src2))
    hit2 = run_logic_page_check(src2, assessment_year=2026, mode="full", domain_id="power_supply")
    notes = [f for f in hit2["findings"] if f["kind"] == "logic" and "合计" in f["note"]]
    assert notes
    assert "10" in notes[0]["note"] and "13" in notes[0]["note"]


def test_llm_gate_off_returns_llm_unused(tmp_path):
    """LOGIC_LLM_GATE=0 时不调用模型，llm_used=False。"""
    src = _save(
        tmp_path,
        [
            ("第3章 设备功能有效性评估", True),
            ("导线松驰，需要安排复测。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    assert hit["llm_used"] is False
    assert hit["llm_finding_count"] == 0


def test_table_sum_check_unit():
    """表格核对函数直接验证：分项 5+8 对合计 10 报错，对合计 13 不报。"""
    from chapters.common.logic_page_check import _table_sum_check

    rows = [
        ["线路", "故障起数", "备注"],
        ["2号线", "5", ""],
        ["3号线", "8", ""],
        ["合计", "10", ""],
    ]
    notes = _table_sum_check(rows, "第8章")
    assert notes and "13" in notes[0]["note"]
    rows_ok = [r[:] for r in rows]
    rows_ok[3][1] = "13"
    assert _table_sum_check(rows_ok, "第8章") == []


def test_index_column_is_not_checked_for_sum():
    """序号列不做合计核对（表头写序号、合计行写行号时不得误报），真实数值列仍报错。"""
    from chapters.common.logic_page_check import _table_sum_check

    # 还原供电报告「生产计划执行情况」表：序号 1~18，合计行序号写 19；
    # 第10、11、16行计划数量为「无数据」，其余 15 行分项和 57615，合计行写 86477（真实错误）。
    plan_full = [
        4119, 6372, 4212, 2538, 3101, 2931, 4006, 3711, 4855,
        None, None, 3948, 3227, 4006, 4217, None, 2660, 3712,
    ]
    rows = [["序号", "线路", "计划数量(项)", "完成数量(项)", "未完成数量(项)", "完成率"]]
    for n in range(1, 19):
        v = plan_full[n - 1]
        p = str(v) if v is not None else "无数据"
        rows.append([str(n), f"{n}号线", p, p, "0", "100%"])
    rows.append(["19", "合计(项)", "86477", "86477", "0", "100%"])
    notes = _table_sum_check(rows, "第7章")
    blob = " ".join(x["note"] for x in notes)
    # 序号列不得报「1+2+…+18=171 与 19 不符」
    assert "171" not in blob and "序号" not in blob
    # 真实的数值列合计错误仍要查出
    assert "计划数量" in blob and "86477" in blob and "57615" in blob

    # 表头缺失时，分项值是连续自然数 1,2,3 也判为序号列，不报错
    bare = [
        ["", "数量"],
        ["1", "5"],
        ["2", "8"],
        ["3", "合计"],
    ]
    assert _table_sum_check(bare, "第7章") == []


def _grade_scrap_units():
    """构造第3章评级表 + 第11章报废清单的 units。"""
    grade_table = {
        "kind": "table",
        "chapter_no": 3,
        "text": "线路 / 区段 / 电力监控设备",
        "rows": [
            ["线路", "区段", "电力监控设备"],
            ["1号线", "正线", "A"],
            ["2号线", "正线", "B"],
            ["3号线", "正线", "C"],
            ["4号线", "正线", "D"],
            ["5号线", "正线", "A"],
            ["6号线", "正线", "B"],
        ],
    }
    map_table = {
        "kind": "table",
        "chapter_no": 3,
        "text": "大类 / 中类名称 / 包含小类设备",
        "rows": [
            ["大类", "中类名称", "包含小类设备"],
            ["供电", "电力监控系统", "中央信号屏、信号屏"],
        ],
    }
    scrap = [
        {"kind": "heading", "chapter_no": 11, "text": "轨道交通1号线", "level": 4},
        {"kind": "paragraph", "chapter_no": 11, "text": "1号线需报废2台中央信号屏，原值共计10万元。"},
        {"kind": "heading", "chapter_no": 11, "text": "轨道交通2号线", "level": 4},
        {"kind": "paragraph", "chapter_no": 11, "text": "2号线需报废2台中央信号屏，原值共计10万元。"},
        {"kind": "heading", "chapter_no": 11, "text": "轨道交通3号线", "level": 4},
        {"kind": "paragraph", "chapter_no": 11, "text": "3号线需报废2台中央信号屏，原值共计10万元。"},
        {"kind": "heading", "chapter_no": 11, "text": "轨道交通4号线", "level": 4},
        {"kind": "paragraph", "chapter_no": 11, "text": "4号线需报废2台中央信号屏，原值共计10万元。"},
        {"kind": "heading", "chapter_no": 11, "text": "轨道交通5号线", "level": 4},
        {"kind": "paragraph", "chapter_no": 11, "text": "5号线需报废2台空调，原值共计2万元。"},
    ]
    return [grade_table, map_table] + scrap


def test_grade_scrap_a_b_c_grades_are_contradictions():
    """评 A、B 或 C 且报废均报矛盾；D 不报；评级表没有的类别（空调）忽略。"""
    from chapters.common.grade_scrap_check import grade_scrap_findings

    rows = grade_scrap_findings(_grade_scrap_units())
    assert len(rows) == 3
    note1 = next(r["note"] for r in rows if "1号线" in r["note"])
    note2 = next(r["note"] for r in rows if "2号线" in r["note"])
    note3 = next(r["note"] for r in rows if "3号线" in r["note"])
    assert all(r["kind"] == "severe" for r in rows)
    assert "电力监控设备" in note1 and "评级为A" in note1 and "中央信号屏" in note1
    assert "电力监控设备" in note2 and "评级为B（状态良，75~90分）" in note2
    assert "电力监控设备" in note3 and "评级为C（状态中，60~75分）" in note3
    # 4号线评 D 不报；5号线报废的空调不在评级表，忽略
    assert not any(("4号线" in r["note"]) or ("5号线" in r["note"]) for r in rows)


def test_grade_scrap_only_d_no_findings():
    """评级表仅 D（差，60 分以下）时不产生矛盾。"""
    from chapters.common.grade_scrap_check import grade_scrap_findings

    units = _grade_scrap_units()
    for r in units[0]["rows"][1:]:
        r[2] = "D"  # 所有线路评级全部改为 D
    assert grade_scrap_findings(units) == []


def test_chapter_titles_without_number_are_recognized(tmp_path):
    """Word 自动编号标题（解析后只有章名）也能识别章号，并跑通评级×报废核对。"""
    src = tmp_path / "no_number.docx"
    doc = Document()
    # 目录区：低级标题中的章名不得当作正文第3章
    doc.add_heading("评估的内容", level=3)
    doc.add_heading("设备功能有效性评估", level=4)
    doc.add_heading("退运报废倾向性评估", level=4)
    # 正文：一级标题（无编号）
    doc.add_heading("设备功能有效性评估", level=1)
    t = doc.add_table(rows=8, cols=3)
    t.rows[0].cells[0].text = "线路"
    t.rows[0].cells[1].text = "区段"
    t.rows[0].cells[2].text = "电力监控设备"
    grades = ["A", "B", "C", "A", "B", "C", "D"]
    for n, g in enumerate(grades, start=1):
        t.rows[n].cells[0].text = f"{n}号线"
        t.rows[n].cells[1].text = "正线"
        t.rows[n].cells[2].text = g
    doc.add_heading("退运报废倾向性评估", level=1)
    doc.add_heading("轨道交通1号线", level=3)
    doc.add_paragraph("1号线需报废2台中央信号屏，原值共计10万元。")
    doc.save(str(src))
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="power_supply")
    severe = [f for f in hit["findings"] if f["kind"] == "severe"]
    assert len(severe) == 1
    assert "第3章" in severe[0]["section"] and "第11章" in severe[0]["section"]


def test_chapter_title_with_status_note_still_locates_chapter():
    """章名后带「（待更新）」时仍是第4章，不能整章并进第3章。"""
    from chapters.common.logic_page_check import _blocks_to_units, _select_units
    from parsers.document_model import Block, DocumentModel

    doc = DocumentModel(
        source_name="供电汇总.docx",
        source_path="a.docx",
        suffix=".docx",
        blocks=[
            Block(type="heading", text="设备功能有效性评估", level=1),
            Block(type="paragraph", text="3号线变压器评为A级。"),
            Block(type="heading", text="运营契合满足度评估（待更新）", level=1),
            Block(type="paragraph", text="4号线故障249起，其中接触网故障。"),
            Block(type="paragraph", text="管理体系合规性评估"),
            Block(type="paragraph", text="本章记录2026年的管理体系情况。"),
        ],
    )
    units = _blocks_to_units(doc)
    picked, msg, _found, _missing = _select_units(units, mode="chapter", chapter_id="ch4")
    assert msg == ""
    blob = " ".join(u["text"] for u in picked)
    assert "249起" in blob
    assert "评为A级" not in blob
    picked5, msg5, _f5, _m5 = _select_units(units, mode="chapter", chapter_id="ch5")
    assert msg5 == ""
    assert any("管理体系情况" in u["text"] for u in picked5)


def test_long_title_note_counts_for_both_domains_and_both_modes():
    """触网长备注、供电短备注：全文和按章都把该章正文算进去，不把「参考第N章」的句子当成章标题。"""
    from chapters.common.logic_page_check import _blocks_to_units, _select_units
    from parsers.document_model import Block, DocumentModel

    note = "运营契合满足度评估（设备体量，章节内结构调整待更新）"
    doc = DocumentModel(
        source_name="触网汇总.doc",
        source_path="oh.doc",
        suffix=".doc",
        blocks=[
            Block(type="heading", text="设备功能有效性评估", level=1),
            Block(type="paragraph", text="接触线评为A级。"),
            Block(type="heading", text=note, level=1),
            Block(type="paragraph", text="4号线接触网故障22起。"),
            Block(type="heading", text="运维表现健康度评估", level=1),
            Block(type="heading", text="风险隐患闭环度，参考第八章评估内容，小结如下", level=3),
            Block(type="paragraph", text="此处只是小结，不是第8章正文。"),
            Block(type="heading", text="使用环境符合性评估", level=1),
            Block(type="paragraph", text="隧道漏水已加装护套。"),
        ],
    )
    units = _blocks_to_units(doc)
    for domain in ("power_supply", "overhead"):
        chapter_picked, msg, _cf, _cm = _select_units(units, mode="chapter", chapter_id="ch4")
        assert msg == "", domain
        chapter_blob = " ".join(u["text"] for u in chapter_picked)
        assert "22起" in chapter_blob
        assert "评为A级" not in chapter_blob
        assert "不是第8章正文" not in chapter_blob
        full_picked, full_msg, full_found, full_missing = _select_units(units, mode="full", chapter_id="ch4")
        # 只识别到第3、4、7、10章，其余缺失时给出缺章警告，防止漏章不自知
        assert {c["no"] for c in full_found} == {3, 4, 7, 10}
        assert set(full_missing) == {5, 6, 8, 9, 11}
        assert "未识别" in full_msg
        full_ch4 = " ".join(u["text"] for u in full_picked if u.get("chapter_no") == 4)
        assert "22起" in full_ch4
        assert "评为A级" not in full_ch4
    picked8, msg8, _f8, _m8 = _select_units(units, mode="chapter", chapter_id="ch8")
    assert picked8 == []
    assert "第8章" in msg8


def test_paragraph_style_chapter_title_recognized_in_order(tmp_path):
    """第5章标题误用正文样式（paragraph）时，按章号顺序也能识别。"""
    src = tmp_path / "para_title.docx"
    doc = Document()
    doc.add_heading("设备功能有效性评估", level=1)
    doc.add_paragraph("导线松驰，需要安排复测。")
    doc.add_paragraph("管理体系合规性评估")
    doc.add_paragraph("本章记录2026年的管理体系情况。")
    doc.save(str(src))
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    # 目录式短章名自身不产生问题；错别字归入第3章
    typo = [f for f in hit["findings"] if f["kind"] == "typo"]
    assert typo
    assert "第3章" in typo[0]["section"]
    assert all("管理体系合规性评估" not in f["note"] for f in hit["findings"])


def test_toc_plain_number_entry_not_flagged(tmp_path):
    """「12 总结与建议 132」这类纯数字编号目录行不报标点错误。"""
    src = _save(
        tmp_path,
        [
            ("目录", False),
            ("12 总结与建议 132", False),
            ("第3章 设备功能有效性评估", True),
            ("导线松驰，需要安排复测。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    assert not any("总结与建议" in (f.get("excerpt") or "") for f in hit["findings"])
    assert any(f["kind"] == "typo" for f in hit["findings"])


def test_grade_from_cell_bands():
    """分数按团标区间换算等级：A≥90、B≥75、C≥60、D<60；字母原样；非评级内容为空。"""
    from chapters.common.grade_scrap_check import _grade_from_cell

    for v in ("90", "95", "100", "100.0"):
        assert _grade_from_cell(v) == "A"
    for v in ("75", "89", "89.9"):
        assert _grade_from_cell(v) == "B"
    for v in ("60", "74"):
        assert _grade_from_cell(v) == "C"
    for v in ("0", "59"):
        assert _grade_from_cell(v) == "D"
    assert _grade_from_cell("A") == "A"
    assert _grade_from_cell("d") == "D"
    assert _grade_from_cell("/") == ""
    assert _grade_from_cell("") == ""
    assert _grade_from_cell("101") == ""


def _overhead_grade_units():
    """触网式评级表：设备列为数字分数、线路列为纯数字、含综合状态列与年度对比表。"""
    table_now = {
        "kind": "table",
        "chapter_no": 3,
        "text": "线路 / 区段 / 刚性接触网 / 柔性接触网 / 接触轨 / 隔离开关 / 系统状态",
        "rows": [
            ["线路", "区段", "刚性接触网", "柔性接触网", "接触轨", "隔离开关", "系统状态值S1", "系统状态等级"],
            ["4", "正线", "/", "98", "/", "100", "99", "A"],
            ["1", "北北延伸", "/", "73", "/", "83", "81", "B"],
            ["8", "一期", "98", "98", "/", "87", "95", "A"],
            ["8", "二期", "98", "73", "/", "79", "85", "B"],
        ],
    }
    table_cmp = {
        "kind": "table",
        "chapter_no": 3,
        "text": "线路 / 区段 / 时间 / 刚性接触网 / 柔性接触网 / 接触轨 / 隔离开关 / 系统状态",
        "rows": [
            ["线路", "区段", "时间", "刚性接触网", "柔性接触网", "接触轨", "隔离开关", "系统状态值S1", "系统状态等级"],
            ["4", "正线", "2025", "/", "99", "/", "57", "77", "B"],
            ["4", "正线", "2026", "/", "98", "/", "100", "99", "A"],
        ],
    }
    scrap = [
        {"kind": "heading", "chapter_no": 11, "text": "2026年设备退运更换情况", "level": 2},
        {"kind": "paragraph", "chapter_no": 11, "text": "4号线： 4号线需报废4台隔离开关，原值共计147812.93元。"},
        {"kind": "heading", "chapter_no": 11, "text": "评估小结", "level": 2},
        {"kind": "paragraph", "chapter_no": 11, "text": "退运报废评估显示，涉及4号线4台隔离开关（服役21年）、8号线二期浦江停车场短电缆及隔离开关（服役17年）等。"},
    ]
    return [table_now, table_cmp] + scrap


def test_overhead_numeric_grade_vs_scrap():
    """触网：数字分数换算（100=A）、纯数字线路列、句中线路号、设备名直配列、小结复述不重复。"""
    from chapters.common.grade_scrap_check import grade_scrap_findings

    rows = grade_scrap_findings(_overhead_grade_units(), assessment_year=2026)
    assert len(rows) == 1
    row = rows[0]
    assert row["kind"] == "severe"
    assert "4号线" in row["note"] and "隔离开关" in row["note"] and "正线" in row["note"]
    # 8 号线隔离开关虽为 B，但只在评估小结复述（无台数、无「需报废」汇总句），不抽取、不报
    assert "8号线" not in row["note"]
    # 综合状态列（系统状态等级）不得作为设备类别
    assert "系统状态" not in row["note"]


def test_grade_scrap_year_filter_uses_assessment_year():
    """年度对比表：2025 评 A、2026 降为 D 时，按评估年（2026）不报错。"""
    from chapters.common.grade_scrap_check import grade_scrap_findings

    units = [
        {
            "kind": "table",
            "chapter_no": 3,
            "rows": [
                ["线路", "区段", "时间", "隔离开关"],
                ["1", "正线", "2025", "A"],
                ["1", "正线", "2026", "D"],
                ["2", "正线", "2025", "B"],
                ["2", "正线", "2026", "B"],
                ["3", "正线", "2025", "C"],
                ["3", "正线", "2026", "C"],
            ],
        },
        {"kind": "paragraph", "chapter_no": 11, "text": "1号线需报废2台隔离开关。"},
    ]
    assert grade_scrap_findings(units, assessment_year=2026) == []
    # 未指定评估年时，旧年份的 A 也会被合并（防御性检查）
    assert len(grade_scrap_findings(units)) == 1


def test_scrap_total_detail_dedup_and_semicolon_clauses():
    """总述数量被分述具体设备同类等量覆盖时不重复计数；分号子句的设备一并抽取。"""
    from chapters.common.grade_scrap_check import grade_scrap_findings

    grade = {
        "kind": "table",
        "chapter_no": 3,
        "rows": [
            ["线路", "区段", "变压器设备", "应急电源设备", "配电系统(牵引)", "电力监控设备"],
            ["8号线", "一期", "B", "B", "B", "B"],
            ["8号线", "二期", "B", "B", "B", "B"],
            ["2号线", "正线", "B", "B", "B", "B"],
        ],
    }
    mapping = {
        "kind": "table",
        "chapter_no": 3,
        "rows": [
            ["大类", "中类名称", "包含小类设备"],
            ["供电", "配电系统（牵引）", "整流器、高压开关"],
        ],
    }
    scrap = [
        {"kind": "paragraph", "chapter_no": 11,
         "text": "8号线有2台变压器、1台UPS电源需报废，1台动力变压器，安装在大世界站；1台整流变压器，安装在芦恒路站。"},
        {"kind": "paragraph", "chapter_no": 11,
         "text": "2号线需报废12台供电系统设备，包含1台整流器，于2000年使用；2台高压开关，于2000年使用；7台中央信号屏，于2000年使用。"},
    ]
    rows = grade_scrap_findings([grade, mapping] + scrap, assessment_year=2026)
    note8 = next(r["note"] for r in rows if "8号线" in r["note"] and "变压器" in r["note"])
    assert "动力变压器1台" in note8 and "整流变压器1台" in note8 and "共2台" in note8
    # 总述「2台变压器」已被分述覆盖，不得重复计数
    assert "变压器2台" not in note8
    note2q = next(r["note"] for r in rows if "2号线" in r["note"] and "牵引" in r["note"])
    assert "整流器1台" in note2q and "高压开关2台" in note2q and "共3台" in note2q
    note2m = next(r["note"] for r in rows if "2号线" in r["note"] and "电力监控" in r["note"])
    assert "中央信号屏7台" in note2m


def test_best_grade_per_section_when_tables_disagree():
    """两张评级表同一区段评级不同（一A一B）时取最高，文案不出现同区段既A又B。"""
    from chapters.common.grade_scrap_check import grade_scrap_findings

    rows1 = [["线路", "区段", "电力监控设备"]] + [
        [f"{n}号线", "正线", "B"] for n in range(1, 7)
    ]
    rows1[1][2] = "A"  # 1号线在第一张表为 A
    rows2 = [["线路", "区段", "电力监控设备"]] + [
        [f"{n}号线", "正线", "B"] for n in range(1, 7)
    ]  # 1号线在第二张表为 B
    units = [
        {"kind": "table", "chapter_no": 3, "rows": rows1},
        {"kind": "table", "chapter_no": 3, "rows": rows2},
        {"kind": "paragraph", "chapter_no": 11, "text": "1号线需报废2台中央信号屏。"},
    ]
    rows = grade_scrap_findings(units)
    assert len(rows) == 1
    assert "评级为A" in rows[0]["note"]
    assert "评级为B" not in rows[0]["note"]


def test_llm_drops_consistent_notes():
    """LLM 把「核对一致/无矛盾」的验证说明返回时，过滤掉，不展示为计算错误。"""
    import chapters.common.logic_llm as m

    blob = "1台中央信号屏于2017年10月开始使用，设计使用年限为10年，至今已使用8年。"
    raw = [
        {"kind": "logic", "excerpt": blob, "note": "2017年10月至2026年应为8年，与年份计算一致。"},
        {"kind": "logic", "excerpt": blob, "note": "2004年12月至2026年应为21年，此处与年份计算一致，无矛盾。"},
    ]
    assert m._clean_findings(raw, blob) == []

    # 真正的矛盾（不一致、应为X但原文写Y）保留
    raw_bad = [
        {"kind": "logic", "excerpt": "4号线故障5起，第8章写12起。", "note": "两章数字不一致，相差7起。"},
        {"kind": "logic", "excerpt": "本类设备共8台，明细合计10台。", "note": "总数应为10，但原文写8。"},
    ]
    out = m._clean_findings(raw_bad, "4号线故障5起，第8章写12起。本类设备共8台，明细合计10台。")
    assert len(out) == 2 and all(x["kind"] == "logic" for x in out)


def test_llm_drops_non_issue_conclusions():
    """模型明确给出「不构成矛盾/未发现矛盾/不属于跨章核对/非错别字」结论时丢弃。"""
    import chapters.common.logic_llm as m

    notes = [
        ("logic", "口径不同，不构成同一指标矛盾。"),
        ("logic", "但均在第11章内，不属于跨章核对。"),
        ("logic", "无直接可比性，未发现跨章矛盾。"),
        ("logic", "两处报废清单未说明批次，易造成总数与分项矛盾。"),
        ("typo", "主变电站与主变电所疑似用字不一致，但无确凿依据，不报。"),
        ("typo", "k1应为K1，大小写不规范，但非错别字。"),
        ("language", "缺空格，但属格式问题。"),
    ]
    raw = [
        {"kind": k, "excerpt": f"某设备于二〇一七年十月开始使用至今已八年整第{i}处。", "note": n}
        for i, (k, n) in enumerate(notes)
    ]
    blob = "".join(x["excerpt"] for x in raw)
    assert m._clean_findings(raw, blob) == []

    # 确凿结论（存在明显矛盾、疑似 typo 不在此列）保留
    keep = [
        {"kind": "logic", "excerpt": "2号线与9号线班组明细完全相同。", "note": "不同线路明细完全相同，存在明显矛盾。"},
        {"kind": "typo", "excerpt": "导线松驰需要复测。", "note": "江扬与前文江杨用字不一致，疑为错别字。"},
    ]
    blob2 = "2号线与9号线班组明细完全相同。导线松驰需要复测。"
    out = m._clean_findings(keep, blob2)
    assert len(out) == 2


def test_llm_drops_year_span_arithmetic_and_batch_scrap():
    """年限算术（依赖基准日与取整）与同章多批次报废数量对比，不报计算错误。"""
    import chapters.common.logic_llm as m

    dropped = [
        ("logic", "2004年12月至2026年为22年，与“至今使用21年”不符。"),
        ("logic", "2007年2月至2026年为19年，与“至今使用18年”不符；2017年10月至2026年为9年，与“至今使用8年”不符。"),
        ("logic", "2005年12月开始使用至评估年2026年，实际使用约20年，与文中“至今已使用21年”不一致。"),
        ("logic", "一处为1台，另一处为5台，同一线路同一设备报废数量明显不一致。"),
        ("logic", "与后文“1号线需报废5台车站短路器”对1号线车站短路器报废数量表述不一致。"),
        ("logic", "两处均为2号线设备报废，一处合计6台，另一处合计12台。"),
        ("logic", "户外段4台、隧道段剩余4台与剩余46台合计关系需核对。"),
    ]
    raw = [
        {"kind": k, "excerpt": f"某设备于二〇〇〇年开始使用至今已使用若干年记录第{i}条。", "note": n}
        for i, (k, n) in enumerate(dropped)
    ]
    blob = "".join(x["excerpt"] for x in raw)
    assert m._clean_findings(raw, blob) == []

    # 真正的数字矛盾（故障起数、完成率）保留
    kept = [
        {"kind": "logic", "excerpt": "第4章2026年故障降压系统58项。", "note": "降压58项与第4章另一处71项明显不一致。"},
        {"kind": "logic", "excerpt": "14站完成13站，完成率87%。", "note": "13/14≈92.9%，与文中87%不符。"},
    ]
    blob2 = "第4章2026年故障降压系统58项。14站完成13站，完成率87%。"
    out = m._clean_findings(kept, blob2)
    assert len(out) == 2


def test_llm_single_chapter_findings_get_chapter_no(monkeypatch):
    """单章 LLM 复核结果必须带章号，前端定位显示「第N章」而非「跨章数字核对」。"""
    import chapters.common.logic_llm as m

    monkeypatch.setattr(m, "logic_llm_enabled", lambda: True)
    monkeypatch.setattr(
        m,
        "_chat",
        lambda sys, user: {
            "findings": [
                {"kind": "typo", "excerpt": "导线松驰需要复测安排。", "note": "松驰应为松弛。"},
            ]
        },
    )
    text = "导线松驰需要复测安排。" * 30
    units = [
        {"kind": "heading", "chapter_no": 3, "text": "设备功能有效性评估"},
        {"kind": "paragraph", "chapter_no": 3, "text": text},
        {"kind": "paragraph", "chapter_no": 3, "text": text},
        {"kind": "paragraph", "chapter_no": 3, "text": text},
    ]
    items, used = m.llm_review_units(units, mode="chapter", chapter_id="ch3", year=2026)
    assert used is True
    assert items and items[0]["chapter_no"] == 3


def test_strip_title_note_nested_and_multiple_brackets():
    """末尾备注支持嵌套括号、连续多组、多种括号类型；正文中的括号不剥离。"""
    from chapters.common.logic_page_check import _strip_title_note

    cases = {
        "退运报废倾向性评估（待更新）": "退运报废倾向性评估",
        "退运报废倾向性评估（设备体量（含3.2节），章节结构调整待更新）": "退运报废倾向性评估",
        "直流应急电源系统MTBF（R）（内容与标题不符）": "直流应急电源系统MTBF",
        "设备功能有效性评估【待更新】〔二次确认〕": "设备功能有效性评估",
        "使用环境符合性评估[待更新]": "使用环境符合性评估",
        "管理体系合规性评估（待更新数据 ）": "管理体系合规性评估",
    }
    for raw, expect in cases.items():
        assert _strip_title_note(raw) == expect, raw
    # 未闭合括号不剥离
    assert _strip_title_note("退运报废倾向性评估（待更新") == "退运报废倾向性评估（待更新"
    # 句中括号（不在末尾）不剥离
    assert _strip_title_note("退运报废（见3.2节）倾向性评估") == "退运报废（见3.2节）倾向性评估"


def _resolver_units():
    """3-11 章完整单元（heading 一级标题），章名可带备注。"""
    from chapters.common.logic_page_check import _blocks_to_units
    from parsers.document_model import Block, DocumentModel

    titles = [
        (3, "设备功能有效性评估"),
        (4, "运营契合满足度评估（待更新）"),
        (5, "管理体系合规性评估"),
        (6, "修程修制匹配性评估"),
        (7, "运维表现健康度评估"),
        (8, "风险隐患闭环度评估"),
        (9, "备件物资保障度评估"),
        (10, "使用环境符合性评估"),
        (11, "退运报废倾向性评估（设备体量（含3.2节），章节结构调整待更新）"),
    ]
    blocks = []
    for no, title in titles:
        blocks.append(Block(type="heading", text=title, level=1))
        blocks.append(Block(type="paragraph", text=f"第{no}章正文，全网各线路运行正常。"))
    doc = DocumentModel(source_name="x.docx", source_path="x.docx", suffix=".docx", blocks=blocks)
    return _blocks_to_units(doc)


def test_resolve_chapters_all_nine_with_notes():
    """带末尾备注（含嵌套括号）的 3-11 章全部识别，章节清单完整。"""
    from chapters.common.logic_page_check import resolve_chapters_rules, _chapters_found

    units = _resolver_units()
    chapter_map, missing = resolve_chapters_rules(units)
    assert missing == []
    assert sorted(chapter_map) == list(range(3, 12))
    # 第11章标题含嵌套括号备注，标题清单里是去备注后的章名
    found = _chapters_found(units, chapter_map)
    by_no = {f["no"]: f for f in found}
    assert by_no[11]["title"] == "退运报废倾向性评估"
    assert by_no[4]["title"] == "运营契合满足度评估"
    assert all(f["confidence"] == "high" for f in found)


def test_resolve_chapters_paragraph_style_and_manual_format():
    """正文样式章标题靠核心词+顺序识别；手工排版加粗大字的章标题靠格式信号识别。"""
    from chapters.common.logic_page_check import resolve_chapters_rules

    units = [
        {"i": 0, "kind": "heading", "text": "设备功能有效性评估", "level": 1,
         "style_name": "Heading 1", "font_size": 0.0, "bold": False, "centered": False},
        {"i": 1, "kind": "paragraph", "text": "正文内容，运行正常。", "level": 0,
         "style_name": "Normal", "font_size": 12.0, "bold": False, "centered": False},
        # 第5章：误用正文样式，无编号、无格式
        {"i": 2, "kind": "paragraph", "text": "管理体系合规性评估（待更新）", "level": 0,
         "style_name": "Normal", "font_size": 12.0, "bold": False, "centered": False},
        {"i": 3, "kind": "paragraph", "text": "本章说明合规情况。", "level": 0,
         "style_name": "Normal", "font_size": 12.0, "bold": False, "centered": False},
        # 第11章：手工排版，写成"11 设备退出与处置评估（待更新）"（无核心词，靠编号+格式）
        {"i": 4, "kind": "paragraph", "text": "11 设备退出与处置评估（待更新）", "level": 0,
         "style_name": "Normal", "font_size": 16.0, "bold": True, "centered": True},
        {"i": 5, "kind": "paragraph", "text": "13号线设备报废。", "level": 0,
         "style_name": "Normal", "font_size": 12.0, "bold": False, "centered": False},
    ]
    chapter_map, missing = resolve_chapters_rules(units)
    assert 5 in chapter_map and chapter_map[5]["unit_idx"] == 2
    assert 11 in chapter_map and chapter_map[11]["unit_idx"] == 4
    assert "字体格式" in chapter_map[11]["via"]


def test_run_logic_page_check_returns_chapter_lists(tmp_path):
    """检测结果含 chapters_found/chapters_missing；缺第8章时 missing 含 8。"""
    src = _save(
        tmp_path,
        [
            ("第3章 设备功能有效性评估", True),
            ("导线松驰，需要安排复测。", False),
            ("第11章 退运报废倾向性评估", True),
            ("13号线设备报废说明。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    found_nos = {c["no"] for c in hit["chapters_found"]}
    assert {3, 11} <= found_nos
    assert 8 in hit["chapters_missing"]
    assert "未识别" in hit["message"]

    chapter_hit = run_logic_page_check(
        src, assessment_year=2026, mode="chapter", chapter_id="ch8", domain_id="overhead"
    )
    assert "第8章" in chapter_hit["message"]


def test_llm_fallback_fills_missing_chapter(monkeypatch):
    """规则缺章且 LLM 可用时，LLM 从候选短行补位（低置信），核心词冲突/位置错误的不采纳。"""
    import chapters.common.logic_llm as m
    from chapters.common.logic_page_check import resolve_chapters

    units = _resolver_units()
    # 删掉第11章标题，换成语义近似但无核心词的标题
    units = [u for u in units if u["text"] != "退运报废倾向性评估（设备体量（含3.2节），章节结构调整待更新）"]
    units.append({"i": 99, "kind": "heading", "text": "设备退出与报废处置评估（待更新）", "level": 1,
                  "style_name": "Heading 1", "font_size": 0.0, "bold": False, "centered": False})
    units.append({"i": 100, "kind": "paragraph", "text": "13号线设备报废。", "level": 0,
                  "style_name": "Normal", "font_size": 12.0, "bold": False, "centered": False})

    monkeypatch.setattr(m, "logic_llm_enabled", lambda: True)
    monkeypatch.setattr(m, "_chat", lambda sys, user: {"assign": [{"no": 11, "cand": 0}]})
    chapter_map, missing = resolve_chapters(units)
    assert 11 in chapter_map
    assert chapter_map[11]["source"] == "llm"
    assert missing == []

    # 核心词冲突：候选含第3章核心词，不得判给第11章
    monkeypatch.setattr(m, "_chat", lambda sys, user: {"assign": [{"no": 11, "cand": 0}]})
    out = m.llm_resolve_chapters(
        units,
        {n: {"no": n, "unit_idx": i * 2} for i, n in enumerate(range(3, 11))},
        [{"unit_idx": 0, "text": "设备功能有效性评估"}],
    )
    assert out == {}


def test_number_name_conflict_candidate_rejected():
    """编号与核心词指向不同章时，该候选作废，不产生错位归属。"""
    from chapters.common.logic_page_check import resolve_chapters_rules

    units = [
        {"i": 0, "kind": "heading", "text": "3 设备功能有效性评估", "level": 1,
         "style_name": "Heading 1", "font_size": 0.0, "bold": False, "centered": False},
        {"i": 1, "kind": "paragraph", "text": "正文。", "level": 0,
         "style_name": "Normal", "font_size": 12.0, "bold": False, "centered": False},
        # 冲突行：编号写 4，核心词却是第11章
        {"i": 2, "kind": "heading", "text": "4 退运报废倾向性评估", "level": 1,
         "style_name": "Heading 1", "font_size": 0.0, "bold": False, "centered": False},
        {"i": 3, "kind": "paragraph", "text": "正文。", "level": 0,
         "style_name": "Normal", "font_size": 12.0, "bold": False, "centered": False},
    ]
    chapter_map, missing = resolve_chapters_rules(units)
    assert 3 in chapter_map
    # 冲突行既不能给第4章也不能给第11章
    assert chapter_map.get(3, {}).get("unit_idx") == 0
    assert 4 in missing and 11 in missing

def test_punct_non_sentence_exempted(tmp_path):
    """逻辑检测页：培训/规程清单、图表数据公式、表题图题、括号备注不要求句末标点。"""
    src = _save(
        tmp_path,
        [
            ("第4章 运营契合满意度评估", True),
            ("1、2025.5《400V开关继电保护装置整定值查阅》培训", False),
            ("⑤应急电源/UPS/蓄电池相关规程", False),
            ("故障设备趋势=-14.0%", False),
            ("当年设备故障值 53； 前三年设备故障均值 12 11 8", False),
            ("表格10-2  不同线路弓架次数据（柔性）", False),
            ("10号线暂无退运报废设备。（电力监控设备处于D）", False),
            ("3号线正线触网子系统和设备状态是 C", False),
            ("6号线柔性接触网49.98条公里，刚性接触网44.3条公里", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="overhead")
    punct = [f for f in hit["findings"] if f["kind"] == "punct"]
    assert punct == [], [f["note"] for f in punct]


def test_punct_hard_error_and_full_sentence_still_flagged(tmp_path):
    """完整句缺句号、分号紧跟句号硬伤在逻辑检测页仍报。"""
    src = _save(
        tmp_path,
        [
            ("第4章 运营契合满意度评估", True),
            ("应急电源故障主要集中在UPS配电柜触摸屏故障及蓄电池电池开关脱扣", False),
            ("直流屏PLC已停产，升级产品可替代，备件充足；。", False),
        ],
    )
    hit = run_logic_page_check(src, assessment_year=2026, mode="full", domain_id="power_supply")
    punct = [f for f in hit["findings"] if f["kind"] == "punct"]
    assert len(punct) >= 2, punct


def test_report_generation_punct_rule_not_loosened():
    """共用的 is_punct_issue 未放宽：报告生成路径仍按原口径（豁免只在逻辑检测层）。"""
    from chapters.common.word import is_punct_issue

    assert is_punct_issue("1、2025.5《400V开关继电保护装置整定值查阅》培训")
    assert is_punct_issue("表格10-2  不同线路弓架次数据（柔性）")
    assert not is_punct_issue("2026年共发生故障249起。")

def test_llm_punct_skippable():
    """LLM 标点提示过滤：非句子/口径矛盾/引用原文跳过；半角与硬伤保留。"""
    from chapters.common.logic_page_check import _llm_punct_skippable

    # 应跳过：非句子、引用、合法分号、条款、图表引出
    assert _llm_punct_skippable(
        "1、2025.5《400V开关继电保护装置整定值查阅》培训", "末句缺句号")
    assert _llm_punct_skippable(
        "16号线没有C和D状态的，维持现有管控措施；", "分号后无后续内容，标点使用不当")
    assert _llm_punct_skippable(
        "实际流程如图5-5所示：", "句末使用冒号，但后文为图题")
    assert _llm_punct_skippable(
        "在行业标准CJJ/T 288-2018《标准》中 3.3 规定：", "“中”与“3.3”之间有多余空格")
    assert _llm_punct_skippable(
        "17号线没有C和D状态的，维持现有管控措施。", "句末标点与上句分号不匹配")
    assert _llm_punct_skippable(
        "避雷器或架空地线已安装在平腕臂底座上方1-1.5m处。", "该句为规程条款内容，句末不应使用句号")
    # 应保留：半角标点、连续标点硬伤、完整句标点问题
    assert not _llm_punct_skippable(
        "于2000年12月开始使用,设计使用年限10年。", "句中使用半角逗号，应为全角逗号")
    assert not _llm_punct_skippable(
        "备件充足；。", "分号后面又紧跟句号")
    assert not _llm_punct_skippable(
        "特种设备，消防，防雷，高空作业规程如表5-1。", "并列词语之间应使用顿号，此处用逗号")


def test_llm_typo_segment_name_not_flagged():
    """北北延伸、西西延伸等合法区段名被 LLM 当多字/叠字时过滤；真实错别字保留。"""
    from chapters.common.logic_page_check import _llm_typo_skippable

    # 合法方位延伸名 + 多字/叠字指控：丢弃
    assert _llm_typo_skippable(
        "1号线北北延伸接触网状态为B。", "北北延伸多写了一个北字，应为北延伸")
    assert _llm_typo_skippable(
        "2号线西西延伸柔性接触网纳入跟踪。", "西西延伸重复了一个西字")
    assert _llm_typo_skippable(
        "9号线三期东延伸已开通。", "三期东延伸疑似多字")
    # 真实错别字：保留
    assert not _llm_typo_skippable(
        "接触网驰度偏大。", "驰度应为弛度")
    # 同段含合法区段名，但 LLM 指控的是另一个真实错别字（非方位字）：保留
    assert not _llm_typo_skippable(
        "1号线北北延伸接触网驰度偏大。", "驰度应为弛度")
    # 出现了非法方位词（如东东北延伸），即便指控多字也保留，交人工核对
    assert not _llm_typo_skippable(
        "1号线东东北延伸。", "东东北延伸多写了一个东字")




def test_llm_false_positive_self_report_and_format():
    """LLM 自我否决/纯格式/模糊诊断/自算趋势四类提示过滤；真实错误保留。"""
    from chapters.common.logic_llm import _is_llm_false_positive

    # 自我否决：算了一通后说“不报告/不在核对范围”
    assert _is_llm_false_positive(
        "趋势计算(143-207)/207≈-30.9%，与-16.4%不一致，但趋势计算规则不在本次核对范围，不报告。",
        "language")
    assert _is_llm_false_positive("该数字差异不在本次核对范围，不报告。", "logic")
    # 纯排版/分段/标题格式
    assert _is_llm_false_positive(
        "“北延伸：”紧接在上一段句末之后，未另起段落，标题与正文混排，格式/语句衔接不当。",
        "language")
    assert _is_llm_false_positive("该区段标题没有单独换行、未分段，排版不规范。", "punct")
    # 二选一的不确定诊断
    assert _is_llm_false_positive(
        "“已对需石龙路停车场数量进行梳理”缺“改造”或“的”，语句不通顺。", "language")
    # 自算趋势公式（原文未给公式）
    assert _is_llm_false_positive(
        "趋势(242-207)/207≈16.9%，与原文-16.4%不一致。", "logic")
    # 占比四舍五入仅 0.0X 个百分点的舍入噪音
    assert _is_llm_false_positive(
        "174/189=92.06%，15/189=7.94%，原文写7.93%，存在0.01个百分点的舍入差异。", "logic")
    # 模型自行假设的推算（非原文数字关系）
    assert _is_llm_false_positive(
        "若月月练和月度培训各12次共24次×74人=1776人次，与888不符。", "logic")
    # 真实错误保留
    assert not _is_llm_false_positive("驰度应为弛度。", "typo")
    assert not _is_llm_false_positive(
        "“共10台，其中4台、3台、2台”，分项相加为9台，与总数10台不符。", "logic")
    # 唯一确定的缺字诊断保留（不被“模糊”规则误杀）
    assert not _is_llm_false_positive(
        "“已对需石龙路停车场”缺少“改造”二字，应为“已对需改造石龙路停车场”。", "language")


def test_chunk_blob_splits_on_sentence_boundary():
    """长章按句子边界分块：不超上限、不切词、内容完整、短文本单块。"""
    from chapters.common.logic_llm import _chunk_blob

    short = "第一段内容。\n第二段内容。"
    assert _chunk_blob(short, 6000) == [short]

    paras = "\n".join(f"第{i}段：这是一段足够长的正文，包含完整句子。" for i in range(400))
    chunks = _chunk_blob(paras, 1000)
    assert len(chunks) > 1
    assert all(len(c) <= 1000 for c in chunks)
    assert "".join("".join(chunks).split()) == "".join(paras.split())
    assert all(c.rstrip().endswith("。") for c in chunks)

    # 单段超长：按句号切，关键句“更新改造项目”保持完整，不被切成“更新改造项”
    long_para = "前缀说明。" * 300 + "应急电源设备更新改造项目，项目2025年开始，2026年全部完成。"
    ch = _chunk_blob(long_para, 1000)
    assert all(len(c) <= 1000 for c in ch)
    assert any("更新改造项目，项目2025年开始" in c for c in ch)


def test_llm_language_filters_punct_lists_enum():
    """language 通道：清单条目标点/句末标点归 punct/列举序号不缺量词，均跳过；正文问题保留。"""
    from chapters.common.logic_page_check import _llm_language_skippable, _llm_enum_skippable

    # 破折号/培训清单条目上的句末标点、缺分隔
    assert _llm_language_skippable("——修改了系统非正常运行方式", "句末缺标点，应为「——修改了系统非正常运行方式。」")
    assert _llm_language_skippable("112025.6《变电设备运维管理培训》", "序号11与日期2025.6之间缺分隔，应为11、2025.6")
    # 正文句末标点问题归 punct 通道
    assert _llm_language_skippable("存在明显的环境干扰因素", "「存在明显的环境干扰因素」句末缺标点")
    # 正文句中缺标点、句子成分残缺保留
    assert not _llm_language_skippable("跟随所投入使用", "“使用”与“跟随所”之间缺少标点，分句粘连，应加“；”")
    assert not _llm_language_skippable("由于零部件的失效率", "句子成分残缺，缺谓语，应补出谓语")

    # 分号后的列举序号不缺量词
    assert _llm_enum_skippable(
        "110kVGIS开关柜故障1项；4电力监控系统故障16项多为综保时间不一致",
        "“4电力监控系统”缺字，应为“4项电力监控系统”")
    # 无分号列举语境的真实数量缺量词保留
    assert not _llm_enum_skippable(
        "全所共有4电力监控系统投入运行", "“4电力监控系统”缺量词，应为“4个电力监控系统”")
def test_findings_sorted_by_chapter_then_section():
    """完整检测结果按 章→节→严重程度 排列；跨章矛盾归首章；无章号排最后。"""
    from chapters.common.logic_page_check import _sort_findings

    rows = [
        {"kind": "typo", "section": "第7章 7.2 维护", "note": "错别字"},
        {"kind": "severe", "section": "第3章 / 第11章 轨道交通1号线", "note": "跨章矛盾"},
        {"kind": "logic", "section": "第5章 5.1 概述", "note": "计算错"},
        {"kind": "punct", "section": "第3章 3.2 设备", "note": "标点"},
        {"kind": "logic", "section": "第3章 3.10 设备", "note": "计算3.10"},
        {"kind": "logic", "section": "第3章 3.1 维护", "note": "计算3.1"},
        {"kind": "year", "section": "跨章数字核对", "note": "年份"},
    ]
    order = [f['note'] for f in _sort_findings(rows)]
    assert order == [
        "跨章矛盾",  # 第3章首（severe，无节号）
        "计算3.1",
        "标点",
        "计算3.10",  # 数值节号 3.10 在 3.2 之后
        "计算错",  # 第5章
        "错别字",  # 第7章
        "年份",  # 无章号排最后
    ]


def test_findings_same_section_severe_first():
    """同一章同一节内，严重问题排在轻微问题之前。"""
    from chapters.common.logic_page_check import _sort_findings

    rows = [
        {"kind": "punct", "section": "第4章 4.1 概述", "note": "标点"},
        {"kind": "severe", "section": "第4章 4.1 概述", "note": "严重"},
        {"kind": "typo", "section": "第4章 4.1 概述", "note": "错别字"},
    ]
    order = [f['note'] for f in _sort_findings(rows)]
    assert order == ["严重", "错别字", "标点"]


def test_llm_dup_prefix_and_suffix_same_error():
    """同一处错误被截取成短/长两条（前缀或后缀）时合并，保留信息更完整的一条。"""
    from chapters.common.logic_page_check import _llm_dup_index

    # scada：短摘录是长摘录的严格前缀
    short = {"kind": "typo", "note": "「scada」大小写不规范，应改为「SCADA」",
             "excerpt": "已纳入2024年1、6、7、8号线scada大修更新改造", "via": "llm", "section": "第3章"}
    long_ = {"kind": "typo", "note": "「scada」大小写不规范，应改为「SCADA」",
             "excerpt": ("已纳入2024年1、6、7、8号线scada大修更新改造，计划2025年完成，"
                         "计划2站，已完成1站，占比50%，计划2025年9月份完成改造更换。"),
             "via": "llm", "section": "第3章"}
    assert _llm_dup_index(long_, [short]) == 0
    assert _llm_dup_index(short, [long_]) == 0

    # 92.06% 顿号：短摘录是长摘录的严格后缀
    head = ("10号线自2025年5月1日起至2026年4月30日故障总数约为189件，其中配电系统69起，"
            "应急电源故障42起，电力监控系统故障36起，杂散电流设备故障27起，")
    tail = ("四类设备合计故障174起，占总故障数的92.06%、其他各类的故障数量较少，"
            "合计仅15起，占总故障数的7.93%。")
    a = {"kind": "punct", "note": "“92.06%”后应用逗号或分号，此处用顿号连接并列分句。",
         "excerpt": tail, "via": "llm", "section": "第4章"}
    b = {"kind": "punct", "note": "“92.06%”后应为逗号或分号，此处用顿号连接并列分句。",
         "excerpt": head + tail, "via": "llm", "section": "第4章"}
    assert _llm_dup_index(b, [a]) == 0


def test_llm_dup_keeps_template_repeated_real_errors():
    """模板化多处同型错误（两处温室度、两处4月低）摘录不同，不合并。"""
    from chapters.common.logic_page_check import _llm_dup_index

    # 两处「温室度」：一条仅多「我方」前缀，差距不足 8 字
    a = {"kind": "typo", "note": "「温室度」用词不当，应为「温湿度」",
         "excerpt": "我方在变电站内装设智能温室度表计，密切监控所内温湿度", "via": "llm"}
    b = {"kind": "typo", "note": "「温室度」应为「温湿度」",
         "excerpt": "在变电站内装设智能温室度表计，密切监控所内温湿度", "via": "llm"}
    assert _llm_dup_index(b, [a]) == -1

    # 两处「4月低」：完成站数不同，摘录互不为前后缀
    a = {"kind": "language", "note": "「4月低前」应为「4月底前」",
         "excerpt": "4月低前完成全线检查，计划完成33站，已完成10站，完成率100%；", "via": "llm"}
    b = {"kind": "language", "note": "「4月低前」应为「4月底前」",
         "excerpt": "4月低前完成全线检查，计划完成33站，已完成33站，完成率100%；", "via": "llm"}
    assert _llm_dup_index(b, [a]) == -1


def test_llm_dup_rule_vs_llm_same_numeric_conflict():
    """规则已报的 9起≠分项合计7起，LLM 再报同一条时去重（保留规则条目）。"""
    from chapters.common.logic_page_check import _llm_dup_index

    rule = {
        "kind": "logic",
        "note": "故障总数9起与分项合计7起（6+1+0+0）不符",
        "excerpt": ("2025年至今共发生故障9起，其中降压系统故障6项；牵引系统故障1项；"
                    "电力监控系统故障0项；应急电源系统故障0项。其中占比较高的分别是"
                    "降压系统故障、牵引系统故障和电力监控系统故障，降压系统多为400V通讯故障，"
                    "牵引系统多为直流开关测试接触器故障。"),
        "via": "",
        "section": "第4章 各个线路基本情况",
    }
    llm = {
        "kind": "logic",
        "note": "分项相加：6+1+0+0=7项，与总数9起不符，缺少2项。",
        "excerpt": ("2025年至今共发生故障9起，其中降压系统故障6项；牵引系统故障1项；"
                    "电力监控系统故障0项；应急电源系统故障0项。其中占比较高的分别是"
                    "降压系统故障、牵引系统故障和电力监控系统故障。"),
        "via": "llm",
        "section": "第4章",
    }
    assert _llm_dup_index(llm, [rule]) == 0


def test_rounding_noise_filtered_but_real_count_conflict_kept():
    """占比 0.01~0.02 个百分点的舍入噪音过滤；整数数量真错误保留。"""
    from chapters.common.logic_llm import _is_rounding_noise_note, _is_llm_false_positive

    # 占比相加 92.06%+7.93%=99.99% 不等于 100%
    assert _is_rounding_noise_note("92.06%+7.93%=99.99%，不等于100%。")
    # 自算 15/189=7.94% 与原文 7.93% 不符（截断）
    assert _is_rounding_noise_note(
        "174/189=92.06%，15/189=7.94%，原文7.93%与15/189计算不符，且92.06%+7.93%=99.99%不等于100%。")
    # 完整走一遍误报过滤
    assert _is_llm_false_positive(
        "四类分项69+42+36+27=174与合计174一致；174+15=189与总数189一致；但174/189=92.06%，"
        "15/189=7.94%，原文7.93%与15/189计算不符，且92.06%+7.93%=99.99%不等于100%。", "logic")
    # 整数数量真错误（9起 vs 分项7起）不受保护，不被过滤
    assert not _is_rounding_noise_note("分项相加：6+1+0+0=7项，与总数9起不符，缺少2项。")
    assert not _is_llm_false_positive("分项相加：6+1+0+0=7项，与总数9起不符，缺少2项。", "logic")


def test_on_progress_reports_ordered_stages_to_100(tmp_path):
    """on_progress 给真实阶段进度：单调非减、落在 0~100、最终到 100，且不改变检测结果。"""
    src = _save(
        tmp_path,
        [
            ("第3章 设备功能有效性评估", True),
            ("我们评估了A、B设备共150台，其中A100台，B20台。", False),
            ("第4章 运营契合满意度评估", True),
            ("统计截止到2024年。", False),
        ],
    )
    stages: list[tuple[int, str]] = []
    hit = run_logic_page_check(
        src,
        assessment_year=2026,
        mode="full",
        domain_id="power_supply",
        on_progress=lambda pct, label: stages.append((int(pct), str(label))),
    )
    assert stages, "检测过程应至少回调一次进度"
    percents = [p for p, _ in stages]
    assert percents == sorted(percents), "进度百分比必须单调非减"
    assert min(percents) >= 0 and max(percents) <= 100
    assert percents[-1] == 100
    assert all(label for _, label in stages)
    # 回调仅用于展示，不得改变检测条目
    baseline = run_logic_page_check(
        src, assessment_year=2026, mode="full", domain_id="power_supply"
    )
    assert [f.get("note") for f in hit["findings"]] == [
        f.get("note") for f in baseline["findings"]
    ]