# -*- coding: utf-8 -*-
"""表格竖合并：材料有 vMerge 则原样成文；推断兜底只合分类列。"""
from docx import Document

from chapters.common.word import add_table, infer_vmerges, new_report_document
from chapters.power.extract import _table_item
from parsers.document_model import Block


def test_infer_vmerges_specialty_not_numbers():
    rows = [
        ["专业", "等级", "2025年", "2026年"],
        ["触网专业", "一级", "3", "3"],
        ["触网专业", "二级", "3", "3"],
        ["触网专业", "三级", "0", "0"],
        ["触网专业", "作业指导书", "5", "6"],
        ["触网专业", "总计", "11", "12"],
    ]
    merges = infer_vmerges(rows)
    assert merges == [(0, 1, 5)]  # 只并「专业」列，不并年份里的「3」


def test_infer_vmerges_skips_material_name_column():
    """同文「综合保护继电器」不能因推断被竖并；大类/中类可以。"""
    rows = [
        ["序号", "设备大类", "设备中类", "设备小类", "物料名称", "规格型号", "计量单位", "安全库存数量"],
        ["1", "变电", "综保装置", "10kV", "综合保护继电器", "P121", "台", "1"],
        ["2", "变电", "综保装置", "10kV", "综合保护继电器", "F650", "台", "1"],
        ["3", "变电", "综保装置", "1500V", "综合保护继电器", "7SJ686+", "台", "2"],
    ]
    merges = infer_vmerges(rows)
    cols = {m[0] for m in merges}
    assert 1 in cols and 2 in cols
    assert 4 not in cols  # 物料名称
    assert 5 not in cols  # 规格型号
    assert 0 not in cols  # 序号


def test_add_table_applies_inferred_vmerge(tmp_path):
    rows = [
        ["专业", "等级", "2026年"],
        ["触网专业", "一级", "3"],
        ["触网专业", "二级", "3"],
        ["触网专业", "总计", "6"],
    ]
    doc = new_report_document()
    add_table(doc, rows)
    out = tmp_path / "vmerge.docx"
    doc.save(str(out))
    loaded = Document(str(out))
    table = loaded.tables[0]
    xml = table._tbl.xml
    assert "vMerge" in xml or "w:vMerge" in xml
    assert len(table.rows) == 4
    assert "触网专业" in table.cell(1, 0).text
    assert table.cell(1, 1).text.strip() == "一级"
    assert table.cell(2, 1).text.strip() == "二级"


def test_add_table_respects_explicit_empty_vmerge(tmp_path):
    """材料无竖合并时传 []，不要再推断把同文列并上。"""
    rows = [
        ["专业", "等级"],
        ["触网专业", "一级"],
        ["触网专业", "二级"],
    ]
    doc = new_report_document()
    add_table(doc, rows, vmerge=[])
    out = tmp_path / "no_vmerge.docx"
    doc.save(str(out))
    xml = Document(str(out)).tables[0]._tbl.xml
    assert "vMerge" not in xml and "w:vMerge" not in xml


def test_table_item_carries_source_vmerge():
    block = Block(
        type="table",
        rows=[["设备大类", "物料名称"], ["变电", "A"], ["变电", "B"]],
        vmerge=[(0, 1, 2)],
    )
    item = _table_item(block)
    assert item["vmerge"] == [[0, 1, 2]]


def test_infer_vmerges_line_and_segment_nested():
    """表3-9/3-10：同线路竖并；同区段只在同一线路内并，跨线的「正线」不合。"""
    rows = [
        ["线路", "区段", "刚性接触网", "系统状态等级"],
        ["1号线", "北延伸", "/", "B"],
        ["1号线", "南延伸", "/", "B"],
        ["1号线", "正线", "/", "A"],
        ["2号线", "东延伸", "61", "C"],
        ["2号线", "正线", "/", "B"],
        ["3号线", "正线", "/", "B"],
    ]
    merges = infer_vmerges(rows)
    assert (0, 1, 3) in merges  # 1号线
    assert (0, 4, 5) in merges  # 2号线
    assert (1, 1, 3) not in merges  # 区段不同
    assert not any(m[0] == 1 and m[1] <= 5 and m[2] >= 6 for m in merges)  # 2/3号线正线不合
    assert not any(m[0] == 3 for m in merges)  # 等级 B 不合


def test_infer_vmerges_year_compare_pairs_segment():
    """表3-10：同一线路同一区段的 2025/2026 两行合并线路和区段。"""
    rows = [
        ["线路", "区段", "时间", "柔性接触网"],
        ["1号线", "北延伸", "2025", "96"],
        ["1号线", "北延伸", "2026", "98"],
        ["1号线", "正线", "2025", "64"],
        ["1号线", "正线", "2026", "94"],
        ["2号线", "东延伸", "2025", "81"],
        ["2号线", "东延伸", "2026", "58"],
    ]
    merges = infer_vmerges(rows)
    assert (0, 1, 4) in merges  # 1号线四行
    assert (1, 1, 2) in merges  # 北延伸两年
    assert (1, 3, 4) in merges  # 正线两年
    assert (0, 5, 6) in merges  # 2号线
    assert (1, 5, 6) in merges  # 东延伸两年
    assert not any(m[0] == 2 for m in merges)  # 时间列不合


def test_force_rebuild_table_applies_inferred_line_vmerge(tmp_path):
    from chapters.common.table_copy import write_table_item
    from chapters.common.word import new_report_document

    rows = [
        ["线路", "区段", "柔性接触网"],
        ["1号线", "北延伸", "73"],
        ["1号线", "正线", "94"],
        ["2号线", "东延伸", "58"],
    ]
    doc = new_report_document()
    assert write_table_item(doc, {"kind": "table", "rows": rows, "force_rebuild": True})
    xml = doc.tables[0]._tbl.xml
    assert "vMerge" in xml or "w:vMerge" in xml
    out = tmp_path / "t39.docx"
    doc.save(str(out))
    loaded = Document(str(out))
    assert "1号线" in loaded.tables[0].cell(1, 0).text
    assert loaded.tables[0].cell(1, 1).text.strip() == "北延伸"
    assert loaded.tables[0].cell(2, 1).text.strip() == "正线"

