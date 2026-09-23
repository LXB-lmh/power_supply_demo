# -*- coding: utf-8 -*-
"""3.1/3.2/3.3.1：认小节名字不认编号；今年材料 → 去年完整报告 → 黄标题不能复用。"""
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from chapters.power.extract import extract_all, match_ch3_prose_key, resolve_ch3_method, resolve_ch3_prose
from chapters.power.style import CH3_METHOD, CH3_METHOD_MISSING, GRADE_TABLE
from chapters.power.write import write_chapter_docx
from parsers.document_model import Block, DocumentModel


def _doc(name: str, blocks: list[Block], path: str = "") -> DocumentModel:
    return DocumentModel(
        source_name=name,
        source_path=path or name,
        suffix=".docx",
        blocks=blocks,
    )


YEAR_BODY = "依据今年材料中的评估方法和内容，对供电专业设施设备评估结果进行归纳整理。"
PRIOR_BODY = "这是去年完整报告里的3.1评估方法和内容原文，应当在今年材料没有时复用。"
STD_BODY = "设备状态分为A、B、C、D四类设备。这是评估标准正文，不能因为标题写成3.1就当成评估方法。"
SCOPE_BODY = "依据分类与编码规范对管辖供电专业设备进行设施设备评估，各系统包含第二层设备如表所示。"


def _method_blocks(body: str) -> list[Block]:
    return [
        Block(type="heading", text="3  设备功能有效性评估", level=1),
        Block(type="heading", text="3.1  评估方法和内容", level=2),
        Block(type="paragraph", text=body),
        Block(type="heading", text="3.2  评估标准", level=2),
        Block(type="paragraph", text="标准正文不应进入3.1"),
        Block(type="heading", text="3.3  供电子系统评估", level=2),
    ]


def test_match_name_not_number():
    assert match_ch3_prose_key("3.1评估标准") == "standard"
    assert match_ch3_prose_key("3.2评估方法和内容") == "method"
    # 单写「评估方法」是第2章导图节，不能当3.1
    assert match_ch3_prose_key("3.1评估方法") is None
    assert match_ch3_prose_key("评估方法") is None
    assert match_ch3_prose_key("3.3.1供电子系统设备评估范围") == "scope"
    assert match_ch3_prose_key("3.1供电子系统设备评估范围") == "scope"
    assert match_ch3_prose_key("3.3供电子系统评估") is None
    # 正文「评估范围包括…」不能当成 3.3.1 标题
    assert match_ch3_prose_key("根据规范对设备进行评估，评估范围包括供电、能源系统") is None


def test_scope_331_reuses_prior_when_year_missing():
    """3.3.1 与 3.1/3.2 同为可复用体例：今年没有则用去年报告。"""
    year = _doc(
        "今年材料.docx",
        [
            Block(type="heading", text="3.1  评估方法和内容", level=2),
            Block(type="paragraph", text=YEAR_BODY),
        ],
    )
    prior = _doc(
        "去年完整报告.docx",
        [
            Block(type="heading", text="3.3.1  供电子系统设备评估范围", level=3),
            Block(type="paragraph", text=SCOPE_BODY),
            Block(
                type="table",
                rows=[["系统", "设备范围"], ["供电", "变电所及附属"]],
            ),
            Block(type="heading", text="3.3.2  各线路供电子系统状态分布", level=3),
        ],
        path="prior/去年.docx",
    )
    hit = resolve_ch3_prose([year], [prior])
    assert hit["method"]["via"] == "year"
    assert hit["scope"]["via"] == "prior"
    assert SCOPE_BODY in (hit["scope"].get("paras") or [])
    assert any(x.get("kind") == "table" for x in hit["scope"].get("flow") or [])


def test_ch2_assessment_method_heading_not_used_as_31():
    """去年报告第2章「评估方法」+图2-2长文，不能压过真正的3.1短段。"""
    short = "依据供电设施设备的多维度评估参数，采用模糊层次分析法，PHM多维度评估，以及加权打分的策略，对供电专业设施设备评估结果进行归纳整理。"
    long_ch2 = (
        "上图2-2是实施方案导图，项目评估的输入包括：分析方法、评估范围与维度、"
        "资产健康九项流程、现场资料、数据图表视频图像分析、状态预测、整改方案、评估结论。"
        + ("九维度说明。" * 40)
    )
    prior = _doc(
        "去年完整报告.docx",
        [
            Block(type="heading", text="2  评估实施方案", level=1),
            Block(type="heading", text="评估方法", level=2),
            Block(type="paragraph", text=long_ch2),
            Block(type="heading", text="3  设备功能有效性评估", level=1),
            Block(type="heading", text="3.1  评估方法和内容", level=2),
            Block(type="paragraph", text=short),
            Block(type="heading", text="3.2  评估标准", level=2),
            Block(type="paragraph", text=STD_BODY),
        ],
        path="prior/去年完整报告.docx",
    )
    hit = resolve_ch3_method([], [prior])
    assert hit["paras"] == [short]
    assert "图2-2" not in "".join(hit["paras"])
    assert hit["via"] == "prior"


def test_scope_formula_artifacts_become_formula_3_1():
    """3.3.1 没有可回拷公式时，退回体例 formula_3_1。"""
    from chapters.power.style import CH3_FORMULA_LEAD, SCOPE_LEAD

    prior = _doc(
        "去年完整报告.docx",
        [
            Block(type="heading", text="3  设备功能有效性评估", level=1),
            Block(type="heading", text="3.3.1  供电子系统设备评估范围", level=3),
            Block(type="paragraph", text=SCOPE_LEAD),
            Block(type="paragraph", text=CH3_FORMULA_LEAD),
            Block(type="paragraph", text="（3-1）"),
            Block(type="drawing", text=""),
            Block(type="paragraph", text="其中："),
            Block(type="drawing", text=""),
            Block(type="drawing", text=""),
            Block(type="drawing", text=""),
            Block(type="heading", text="3.3.2  各线路供电子系统状态分布", level=3),
        ],
    )
    hit = resolve_ch3_prose([], [prior])["scope"]
    kinds = [x.get("kind") for x in hit["flow"]]
    assert "formula_3_1" in kinds
    assert "（3-1）" not in hit["paras"]
    assert "其中：" not in hit["paras"]
    assert CH3_FORMULA_LEAD in hit["paras"]


def test_scope_keeps_source_formula_blocks():
    """源稿有 formula 块时优先保留，成文按 source_index 回拷。"""
    from chapters.power.style import CH3_FORMULA_LEAD, SCOPE_LEAD

    prior = _doc(
        "去年完整报告.docx",
        [
            Block(type="heading", text="3  设备功能有效性评估", level=1),
            Block(type="heading", text="3.3.1  供电子系统设备评估范围", level=3),
            Block(type="paragraph", text=SCOPE_LEAD),
            Block(type="paragraph", text=CH3_FORMULA_LEAD),
            Block(type="formula", text="（3-1）", source_index=230),
            Block(type="paragraph", text="其中："),
            Block(type="formula", text="", source_index=232),
            Block(type="formula", text="", source_index=233),
            Block(type="formula", text="", source_index=234),
            Block(type="heading", text="3.3.2  各线路供电子系统状态分布", level=3),
        ],
        path="prior/去年完整报告.docx",
    )
    hit = resolve_ch3_prose([], [prior])["scope"]
    kinds = [x.get("kind") for x in hit["flow"]]
    assert kinds.count("formula") == 4
    assert "formula_3_1" not in kinds
    assert "其中：" in hit["paras"]


def test_swapped_number_still_follows_section_name():
    doc = _doc(
        "编号写错.docx",
        [
            Block(type="heading", text="3  设备功能有效性评估", level=1),
            Block(type="heading", text="3.1  评估标准", level=2),
            Block(type="paragraph", text=STD_BODY),
            Block(type="table", rows=GRADE_TABLE),
            Block(type="heading", text="3.2  评估方法和内容", level=2),
            Block(type="paragraph", text=YEAR_BODY),
            Block(type="heading", text="3.3  供电子系统评估", level=2),
            Block(type="heading", text="3.3.2  各线路供电子系统状态分布", level=3),
            Block(type="paragraph", text="状态分布正文不应进入评估范围。"),
        ],
    )
    hit = resolve_ch3_prose([doc], [])
    assert hit["method"]["paras"] == [YEAR_BODY]
    assert hit["standard"]["paras"] == [STD_BODY]
    assert any(x.get("kind") == "table" for x in hit["standard"]["flow"])
    assert hit["scope"]["paras"] == []


def test_resolve_prefers_year_material_over_prior():
    year = _doc("设备功能有效性说明.docx", _method_blocks(YEAR_BODY))
    prior = _doc("2025年年度评估报告（供电）.docx", _method_blocks(PRIOR_BODY))
    hit = resolve_ch3_method([year], [prior])
    assert hit["via"] == "year"
    assert hit["paras"] == [YEAR_BODY]
    assert hit["source"] == "设备功能有效性说明.docx"


def test_resolve_reuses_prior_when_year_has_no_method():
    grade = _doc(
        "设备评估结果总表.xlsx",
        [Block(type="heading", text="工作表:供电", level=1)],
    )
    prior = _doc(
        "2025年年度评估报告（供电）.docx",
        [
            Block(type="heading", text="3  设备功能有效性评估", level=1),
            Block(type="heading", text="3.1  评估方法和内容", level=2),
            Block(type="paragraph", text=PRIOR_BODY),
            Block(type="heading", text="3.2  评估标准", level=2),
            Block(type="paragraph", text=STD_BODY),
            Block(type="heading", text="3.3.1  供电子系统设备评估范围", level=3),
            Block(type="paragraph", text=SCOPE_BODY),
        ],
    )
    hit = resolve_ch3_prose([grade], [prior])
    assert hit["method"]["via"] == "prior"
    assert hit["method"]["paras"] == [PRIOR_BODY]
    assert hit["standard"]["via"] == "prior"
    assert hit["standard"]["paras"] == [STD_BODY]
    assert hit["scope"]["via"] == "prior"
    assert hit["scope"]["paras"] == [SCOPE_BODY]


def test_resolve_skips_line_report_and_overhead_annual():
    line = _doc("1号线评估报告.docx", _method_blocks("线路报告里的3.1不能当供电3.1。"))
    overhead = _doc("2025年年度评估报告（接触网）.docx", _method_blocks("接触网年报3.1不能进供电。"))
    hit = resolve_ch3_method([line, overhead], [])
    assert hit["paras"] == []
    assert hit["via"] == ""


def test_extract_all_does_not_mix_prior_into_grade_tables():
    prior = _doc("2025年年度评估报告（供电）.docx", _method_blocks(PRIOR_BODY))
    pack = extract_all([], year=2026, prior_docs=[prior])
    method = pack["ch3"]["method"]
    assert method["via"] == "prior"
    assert method["paras"] == [PRIOR_BODY]
    assert not pack["ch3"]["power_table"]


def test_write_fills_year_method_and_source_bracket(tmp_path):
    out = tmp_path / "ch3.docx"
    write_chapter_docx(
        {
            "year": 2026,
            "ch3": {"method": {"paras": [YEAR_BODY], "source": "方法说明.docx", "via": "year"}},
        },
        out,
        "ch3",
    )
    doc = Document(str(out))
    texts = [p.text.strip() for p in doc.paragraphs]
    heading = next(p for p in doc.paragraphs if p.text.startswith("3.1"))
    assert heading.text == "3.1  评估方法和内容（方法说明.docx）"
    assert not any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in heading.runs)
    assert YEAR_BODY in texts
    assert CH3_METHOD_MISSING not in heading.text


def test_write_yellow_heading_when_nothing_to_reuse(tmp_path):
    out = tmp_path / "ch3.docx"
    write_chapter_docx({"year": 2026, "ch3": {}}, out, "ch3")
    doc = Document(str(out))
    texts = [p.text.strip() for p in doc.paragraphs]
    for prefix in ("3.1", "3.2", "3.3.1"):
        heading = next(p for p in doc.paragraphs if p.text.startswith(prefix))
        assert CH3_METHOD_MISSING in heading.text
        assert any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in heading.runs)
    assert CH3_METHOD not in texts
