# -*- coding: utf-8 -*-
"""评估材料标黄段落：年报同步标黄并加（原材料标黄）。触网/供电成文共用 add_para。"""
import json

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from chapters.common.source_yellow import (
    SOURCE_YELLOW_NOTE,
    apply_source_yellow,
    looks_source_yellow,
    register_current_docs,
    register_text,
)
from chapters.common.word import add_para
from chapters.overhead.write import _write_flow as oh_write_flow
from chapters.power.write import _write_flow as power_write_flow
from parsers.document_model import Block, DocumentModel
from parsers.docx_parser import parse_docx
from parsers.parse_cache import _cache_has_yellow_schema, load_document


def _is_yellow(p) -> bool:
    return any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in p.runs)


def test_parse_docx_detects_highlight_and_shading(tmp_path):
    src = tmp_path / "yellow.docx"
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("本段在材料里用高亮标黄。")
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph("本段没有标黄。")
    shaded = doc.add_paragraph("本段用黄底纹标黄。")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "FFFF00")
    shaded._p.get_or_add_pPr().append(shd)
    doc.save(str(src))

    parsed = parse_docx(src)
    paras = [b for b in parsed.blocks if b.type == "paragraph"]
    assert [b.yellow for b in paras] == [True, False, True]


def test_add_para_marks_registered_source_yellow():
    register_text("维持现有管控措施，加强巡视。")
    doc = Document()
    add_para(doc, "维持现有管控措施，加强巡视。")
    add_para(doc, "这是普通未标黄段落。")
    add_para(doc, "本节暂无材料。", yellow=True)

    body = [p for p in doc.paragraphs if (p.text or "").strip()]
    assert body[0].text == "维持现有管控措施，加强巡视。" + SOURCE_YELLOW_NOTE
    assert _is_yellow(body[0])
    assert body[1].text == "这是普通未标黄段落。"
    assert not _is_yellow(body[1])
    assert SOURCE_YELLOW_NOTE not in body[2].text
    assert _is_yellow(body[2])


def test_list_prefix_rewrite_still_matches():
    register_text("1、专项更换横跨绝缘子。")
    assert looks_source_yellow("3、专项更换横跨绝缘子。")
    text, hit = apply_source_yellow("3、专项更换横跨绝缘子。")
    assert hit
    assert text.endswith(SOURCE_YELLOW_NOTE)


def test_prior_docs_are_not_registered():
    prior = DocumentModel(
        source_name="去年体例.docx",
        source_path="去年体例.docx",
        suffix=".docx",
        blocks=[Block(type="paragraph", text="去年体例里的标黄句不应带进今年。", yellow=True)],
    )
    current = DocumentModel(
        source_name="今年.docx",
        source_path="今年.docx",
        suffix=".docx",
        blocks=[Block(type="paragraph", text="今年材料标黄的管控措施要保留。", yellow=True)],
    )
    register_current_docs([current])
    assert looks_source_yellow("今年材料标黄的管控措施要保留。")
    assert not looks_source_yellow("去年体例里的标黄句不应带进今年。")
    register_current_docs([prior])
    assert looks_source_yellow("去年体例里的标黄句不应带进今年。")
    register_current_docs([current])
    assert not looks_source_yellow("去年体例里的标黄句不应带进今年。")


def test_overhead_and_power_write_flow_mark_source_yellow():
    flow = [
        {"kind": "para", "text": "接触网材料标黄段落应带标记。", "source_yellow": True},
        {"kind": "para", "text": "接触网普通段落。"},
    ]
    oh = Document()
    oh_write_flow(oh, flow)
    oh_paras = [p for p in oh.paragraphs if (p.text or "").strip()]
    assert oh_paras[0].text.endswith(SOURCE_YELLOW_NOTE)
    assert _is_yellow(oh_paras[0])
    assert SOURCE_YELLOW_NOTE not in oh_paras[1].text

    power = Document()
    power_write_flow(
        power,
        [
            {"kind": "para", "text": "供电材料标黄段落应带标记。", "source_yellow": True},
            {"kind": "para", "text": "供电普通段落。"},
        ],
    )
    pw = [p for p in power.paragraphs if (p.text or "").strip()]
    assert pw[0].text.endswith(SOURCE_YELLOW_NOTE)
    assert _is_yellow(pw[0])
    assert SOURCE_YELLOW_NOTE not in pw[1].text


def test_old_parse_cache_without_yellow_is_stale(tmp_path, monkeypatch):
    from parsers import parse_cache as pc

    monkeypatch.setattr(pc, "_DOCS", tmp_path)
    stale = {
        "source_name": "a.docx",
        "source_path": "a.docx",
        "suffix": ".docx",
        "blocks": [{"type": "paragraph", "text": "旧缓存没有 yellow 字段"}],
    }
    (tmp_path / "deadbeef.json").write_text(json.dumps(stale), encoding="utf-8")
    assert _cache_has_yellow_schema(stale) is False
    assert load_document("deadbeef") is None

    fresh = dict(stale)
    fresh["blocks"] = [{"type": "paragraph", "text": "新缓存", "yellow": False}]
    (tmp_path / "cafebabe.json").write_text(json.dumps(fresh), encoding="utf-8")
    assert _cache_has_yellow_schema(fresh) is True
    hit = load_document("cafebabe")
    assert hit is not None
    assert hit.blocks[0].yellow is False
