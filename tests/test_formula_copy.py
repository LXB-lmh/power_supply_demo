# -*- coding: utf-8 -*-
"""源稿 OMML 公式应作为 formula 块抽取，并整段拷进成稿。"""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from chapters.common.drawings import copy_formula_paragraph_from_body_index
from chapters.power.extract import resolve_ch3_prose
from chapters.power.write import write_chapter_docx
from parsers.docx_parser import parse_docx

# 固定夹具：往年供电年报（含 OMML 公式），避免依赖 uploads 目录的运行时状态
_PRIOR_FIXTURE = Path(__file__).resolve().parent / "data" / "prior_formula_2025.docx"

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def test_parser_keeps_omml_as_formula():
    assert _PRIOR_FIXTURE.is_file(), "缺少公式回归夹具 tests/data/prior_formula_2025.docx"
    doc = parse_docx(_PRIOR_FIXTURE)
    formulas = [b for b in doc.blocks if b.type == "formula"]
    assert len(formulas) >= 3
    assert any("（3-1）" in (b.text or "") for b in formulas)


def test_write_copies_omml_formula_into_report(tmp_path):
    assert _PRIOR_FIXTURE.is_file(), "缺少公式回归夹具 tests/data/prior_formula_2025.docx"
    src = parse_docx(_PRIOR_FIXTURE)
    hit = resolve_ch3_prose([], [src])
    scope = hit["scope"]
    assert any(x.get("kind") == "formula" for x in scope.get("flow") or [])
    out = tmp_path / "ch3_formula.docx"
    write_chapter_docx({"year": 2026, "ch3": {"scope": scope}}, out, "ch3")
    dest = Document(str(out))
    body = dest.element.body
    omml = body.findall(f".//{{{M_NS}}}oMath") + body.findall(f".//{{{M_NS}}}oMathPara")
    assert omml, "成稿应保留源稿 OMML 公式"
    # 公式编号段也应带回（可能是嵌入对象）
    texts = [p.text.strip() for p in dest.paragraphs]
    assert any("综合评估表达式" in t for t in texts)
    assert any("其中" in t for t in texts)


def test_copy_formula_paragraph_api():
    assert _PRIOR_FIXTURE.is_file(), "缺少公式回归夹具 tests/data/prior_formula_2025.docx"
    src = parse_docx(_PRIOR_FIXTURE)
    formula = next(b for b in src.blocks if b.type == "formula" and (b.text or "").strip() == "")
    dest = Document()
    assert copy_formula_paragraph_from_body_index(str(_PRIOR_FIXTURE), formula.source_index, dest) is True
    assert dest.element.body.findall(f".//{{{M_NS}}}oMath") or dest.element.body.findall(
        f".//{qn('w:drawing')}"
    )
