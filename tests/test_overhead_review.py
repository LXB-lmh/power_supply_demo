# -*- coding: utf-8 -*-
"""接触网各章网页评估结果：缺材料、逻辑不通、年份错误，并与 Word 黄标对齐。"""
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from chapters.overhead.review import review_chapter
from chapters.overhead.write import write_chapter_docx


def _pack(outline):
    return {"year": 2026, "chapter_name": "设备功能有效性评估", "outline": outline}


def test_empty_section_is_missing_and_summary_stays_missing():
    rows = review_chapter(
        _pack(
            [
                {"depth": 0, "level": 1, "num": "3", "title": "设备功能有效性评估"},
                {"depth": 1, "level": 2, "num": "3.1", "title": "评估方法和内容", "fill": {"empty": True}},
                {"depth": 1, "level": 2, "num": "3.5", "title": "评估小结", "fill": {"empty": True}},
            ]
        ),
        "ch3",
    )
    issues = {(x["section"], x["issue"]) for x in rows}
    assert ("3.1 评估方法和内容", "缺失材料") in issues
    assert ("3.5 评估小结", "缺失材料") in issues


def test_parent_with_children_is_not_missing():
    rows = review_chapter(
        _pack(
            [
                {"depth": 1, "level": 2, "num": "8.1", "title": "设施设备年度突出事件分析"},
                {
                    "depth": 2,
                    "level": 3,
                    "num": "",
                    "title": "8号线联航路接触网设备断裂故障分析",
                    "content_child": True,
                    "fill": {"empty": False, "paras": ["发布0866#抢修令。"], "flow": []},
                },
            ]
        ),
        "ch8",
    )
    assert not any(x["issue"] == "缺失材料" and x["section"].startswith("8.1") for x in rows)


def test_number_mismatch_is_logic_and_heading_yellow(tmp_path):
    text = "本年度共发生故障5起，其中接触网故障2起、隔离开关故障2起。"
    node = {
        "depth": 1,
        "level": 2,
        "num": "4.3",
        "title": "各线路接触网故障趋势分析",
        "fill": {"empty": False, "paras": [text], "flow": [{"kind": "para", "text": text}]},
    }
    rows = review_chapter(_pack([{"depth": 0, "num": "4", "title": "运营契合满足度评估"}, node]), "ch4")
    hit = [x for x in rows if x["issue"] == "逻辑不通"]
    assert hit
    assert "5" in hit[0]["note"] and "4" in hit[0]["note"]
    out = tmp_path / "ch4.docx"
    write_chapter_docx(_pack([{"depth": 0, "level": 1, "num": "4", "title": "运营契合满足度评估"}, node]), out, "ch4")
    doc = Document(str(out))

    def yellow(p):
        return any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in p.runs)

    heads = [p for p in doc.paragraphs if "故障趋势分析" in (p.text or "")]
    bodies = [p for p in doc.paragraphs if "共发生故障5起" in (p.text or "")]
    assert heads and yellow(heads[0])
    assert bodies and yellow(bodies[0])


def test_early_cutoff_year_is_year_error(tmp_path):
    text = "统计截至2024年，接触网故障呈下降趋势。"
    node = {
        "depth": 1,
        "level": 2,
        "num": "4.3",
        "title": "各线路接触网故障趋势分析",
        "fill": {"empty": False, "paras": [text], "flow": [{"kind": "para", "text": text}]},
    }
    rows = review_chapter(_pack([node]), "ch4")
    assert any(x["kind"] == "year" and x["issue"] == "年份错误" and "2024" in x["note"] for x in rows)
    out = tmp_path / "ch4.docx"
    write_chapter_docx(
        {"year": 2026, "chapter_name": "运营契合满足度评估", "outline": [node]},
        out,
        "ch4",
    )
    doc = Document(str(out))

    def yellow(p):
        return any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in p.runs)

    bodies = [p for p in doc.paragraphs if "截至2024年" in (p.text or "")]
    assert bodies and yellow(bodies[0])
