# -*- coding: utf-8 -*-
"""Word 修订标记（w:ins 插入 / w:del 删除）解析测试。

修订模式文档中，python-docx 的 paragraph.text 读不到 w:ins 内的文字，
解析器须按“接受修订后的最终文本”提取：含插入、不含删除。
"""
from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from parsers.docx_parser import accepted_paragraph_text, parse_docx


def _ins_run(text: str):
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "1")
    ins.set(qn("w:author"), "tester")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    ins.append(r)
    return ins


def _del_run(text: str):
    dele = OxmlElement("w:del")
    dele.set(qn("w:id"), "2")
    dele.set(qn("w:author"), "tester")
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText")
    t.text = text
    r.append(t)
    dele.append(r)
    return dele


def _make_doc(tmp: Path) -> Path:
    doc = Document()
    # 章标题：整段文字在插入修订中
    h = doc.add_heading("", level=1)
    h._p.append(_ins_run("风险隐患闭环度评估"))
    # 普通段：保留 run + 删除 run + 插入 run
    p = doc.add_paragraph()
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "保留文字"
    r.append(t)
    p._p.append(r)
    p._p.append(_del_run("应删除的旧文字"))
    p._p.append(_ins_run("新增的内容"))
    # 全段被删除：接受修订后为空，不应产生块
    p2 = doc.add_paragraph()
    p2._p.append(_del_run("整段删除的内容"))
    path = tmp / "revisions.docx"
    doc.save(str(path))
    return path


def test_accepted_text_includes_insert_excludes_delete(tmp_path):
    path = _make_doc(tmp_path)
    doc = Document(str(path))
    paras = doc.paragraphs
    # 标题在 w:ins 中
    assert (paras[0].text or "") == ""  # python-docx 原生读不到
    assert accepted_paragraph_text(paras[0]).strip() == "风险隐患闭环度评估"
    # 混合段：保留 + 插入，无删除
    mixed = accepted_paragraph_text(paras[1])
    assert "保留文字" in mixed
    assert "新增的内容" in mixed
    assert "应删除的旧文字" not in mixed
    # 全删段：接受修订后为空
    assert accepted_paragraph_text(paras[2]).strip() == ""


def test_parse_docx_revision_title_is_heading_and_no_deleted_block(tmp_path):
    path = _make_doc(tmp_path)
    model = parse_docx(path)
    headings = [b for b in model.blocks if b.type == "heading"]
    assert any(b.text.strip() == "风险隐患闭环度评估" for b in headings), headings
    all_text = "\n".join(b.text or "" for b in model.blocks)
    assert "应删除的旧文字" not in all_text
    assert "整段删除的内容" not in all_text
    assert "新增的内容" in all_text
