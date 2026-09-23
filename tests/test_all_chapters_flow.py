# -*- coding: utf-8 -*-
"""各章抽取流口径：表格随正文走；隐患 PDF 碎片拼接；备件优先 docx，库存表不含触网。"""
from pathlib import Path

from docx import Document

from chapters.ch4.power_extract import split_fault_lines
from chapters.power.extract import (
    extract_compliance,
    extract_env,
    extract_hazards,
    extract_plan_table,
    extract_retire,
    extract_stock,
    inspect_pdfs,
)
from chapters.power.write import write_chapter_docx
from parsers.dispatch import parse_file
from parsers.document_model import Block, DocumentModel


def _doc(name: str, blocks: list[Block]) -> DocumentModel:
    return DocumentModel(source_name=name, source_path=name, suffix=Path(name).suffix, blocks=blocks)


def test_extract_keeps_tables_in_chapter_flows():
    env = extract_env(
        [
            _doc(
                "评估材料（使用环境）.docx",
                [
                    Block("heading", "使用环境符合性评估"),
                    Block("paragraph", "变电站环境总体满足运行要求。"),
                    Block("table", rows=[["测点", "温度"], ["站1", "25"]]),
                ],
            )
        ]
    )
    assert any(x.get("kind") == "table" for x in env["flow"])

    retire = extract_retire(
        [
            _doc(
                "11退运材料.docx",
                [
                    Block("heading", "1号线设备退运更换"),
                    Block("paragraph", "本年度完成更换。"),
                    Block("table", rows=[["设备", "数量"], ["开关柜", "2"]]),
                ],
            )
        ]
    )
    assert any(x.get("kind") == "table" for x in retire["sections"][0]["flow"])

    plan = extract_plan_table(
        [
            _doc(
                "生产计划执行.docx",
                [
                    Block("table", rows=[["线路", "计划数量", "完成率"], ["1号线", "10", "100%"]]),
                    Block("drawing", source_index=3),
                ],
            )
        ]
    )
    kinds = [x.get("kind") for x in plan["flow"]]
    assert "table" in kinds
    # 7.2 计划表 flow  deliberately 不含紧随其后的图，避免误拷无关插图
    assert "drawing" not in kinds

    stock = extract_stock(
        [
            _doc(
                "评估材料（备件）.docx",
                [
                    Block("table", rows=[["物料名称", "规格型号", "安全库存"], ["接触器", "A1", "2"]]),
                    Block("table", rows=[["厂商", "型号"], ["甲", "B2"]]),
                ],
            )
        ]
    )
    assert stock["table"][0][0] == "物料名称"
    # 第 9 章只要供电安全库存表；厂商表不进 flow
    assert sum(1 for x in stock["flow"] if x.get("kind") == "table") == 1

    law = extract_compliance(
        [
            _doc(
                "4&5合规性材料(2026).docx",
                [
                    Block("paragraph", "管理体系合规性评估"),
                    Block("paragraph", "法律法规的获取情况"),
                    Block("paragraph", "已建立法律法规获取清单。"),
                    Block("table", rows=[["文件", "文号"], ["安规", "1"]]),
                    Block("paragraph", "企业标准和制度"),
                    Block("paragraph", "制度、标准执行情况检查"),
                    Block("paragraph", "合规性评价"),
                ],
            )
        ]
    )["ch5"]
    assert any(x.get("kind") == "table" for x in law["law_flow"])

    by_line, _ = split_fault_lines(
        [
            _doc(
                "评估材料（故障情况，黄色待更新）.docx",
                [
                    Block("paragraph", "各个线路基本情况"),
                    Block("paragraph", "1号线"),
                    Block("paragraph", "2026年共发生故障10起。"),
                    Block("table", rows=[["类型", "数量"], ["开关", "1"]]),
                ],
            )
        ]
    )
    assert any(x.get("kind") == "table" for x in by_line["1号线"])


def test_write_keeps_table_when_flow_has_no_drawing(tmp_path):
    pack = {
        "year": 2026,
        "ch10": {
            "source": "评估材料（使用环境）.docx",
            "paras": ["变电站环境总体满足运行要求。"],
            "flow": [
                {"kind": "para", "text": "变电站环境总体满足运行要求。"},
                {"kind": "table", "rows": [["测点", "温度"], ["站1", "25"]]},
            ],
        },
    }
    out = tmp_path / "ch10.docx"
    write_chapter_docx(pack, out, "ch10")
    dest = Document(str(out))
    assert dest.tables
    assert dest.tables[0].cell(0, 0).text.strip() == "测点"
    assert dest.tables[0].cell(1, 1).text.strip() == "25"


def test_extract_hazards_merges_pdf_page_fragments():
    data = extract_hazards(
        [
            _doc(
                "供电分公司2026年1月安全隐患排查动态治理情况.pdf",
                [
                    Block("heading", "第1页"),
                    Block("paragraph", "附件"),
                    Block("paragraph", "6 6 6"),
                    Block("paragraph", "截止2026年1月31日，分公司无重大事故隐患，无一"),
                    Block("paragraph", "般A类隐患，完成一般B类隐患1个。"),
                    Block("paragraph", "2 2 2"),
                    Block("paragraph", "本月排查出隐患32项，均已整改完成。"),
                ],
            )
        ]
    )
    texts = [x["text"] for x in data["pdf_paras"]]
    assert any("无重大事故隐患" in t and "一般B类隐患1个" in t for t in texts)  # 跨页半句拼成一句
    assert any("排查出隐患32项" in t for t in texts)
    assert not any(t.strip() in {"附件", "6 6 6"} for t in texts)  # 页眉页码噪声丢掉


def test_stock_prefers_spare_docx_not_rule_pdf():
    spare = _doc(
        "评估材料（备件）.docx",
        [Block("table", rows=[["物料名称", "规格型号", "安全库存"], ["接触器", "A1", "2"]])],
    )
    rule = _doc(
        "附件1：《维保供电安全库存管理规定》.pdf",
        [
            Block("paragraph", "上海申通地铁集团有限公司企业标准"),
            Block("table", rows=[["序号", "大类", "物料名称", "型号", "安全库存配置总数量"], ["1", "触网", "铜银接触线", "CTA-120", "1"]]),
            Block("table", rows=[["序号", "设备中类", "物料名称", "规格型号", "安全库存数量"], ["1", "综保装置", "综合保护继电器", "P121", "2"]]),
        ],
    )
    rule.suffix = ".pdf"
    stock = extract_stock([rule, spare])
    assert stock["source"] == "评估材料（备件）.docx"  # 有备件 docx 时不用规定 PDF
    assert stock["table"][1][0] == "接触器"
    only_rule = extract_stock([rule])
    assert only_rule["table"][1][1] == "综保装置"
    assert "触网" not in "".join(str(c) for r in only_rule["table"] for c in r)  # 仅规定 PDF 也去掉触网行


def test_all_unique_pack_pdfs_are_inspected():
    root = Path(r"F:\材料\2026供电评估_新\评估材料\99-26年评估报告编制材料")
    if not root.exists():
        return
    pdfs = sorted({p.name: p for p in root.rglob("*.pdf") if not p.name.startswith("~")}.values(), key=lambda x: x.name)
    assert pdfs, "编制材料包里应有 PDF"
    docs = [parse_file(p) for p in pdfs]
    checked = inspect_pdfs(docs)
    names = {x["source"] for x in checked}
    assert names == {p.name for p in pdfs}
    topics = {x["source"]: x["topic"] for x in checked}
    assert any(t == "hazard" for t in topics.values())
    assert any(t == "stock_rule" for t in topics.values())
    haz = extract_hazards(docs)
    assert haz["pdf_paras"]
    stock = extract_stock(docs)
    if stock.get("table"):
        assert "触网" not in "".join(str(c) for r in stock["table"][:4] for c in r)  # 供电库存表不含接触网物料
