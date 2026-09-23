# -*- coding: utf-8 -*-
"""接触网成文：paras 与 flow 镜像时不应双写。"""
from docx import Document

from chapters.overhead.write import write_chapter_docx


def test_write_skips_paras_duplicated_in_flow(tmp_path):
    pack = {
        "year": 2026,
        "chapter_name": "设备功能有效性评估",
        "outline": [
            {
                "depth": 1,
                "level": 2,
                "num": "3.1",
                "title": "评估方法和内容",
                "fill": {
                    "empty": False,
                    "source": "测试",
                    "paras": ["段落一。", "段落二。"],
                    "flow": [
                        {"kind": "para", "text": "段落一。"},
                        {"kind": "para", "text": "段落二。"},
                    ],
                },
            }
        ],
    }
    out = tmp_path / "oh.docx"
    write_chapter_docx(pack, out, "ch3")
    texts = [p.text for p in Document(str(out)).paragraphs if (p.text or "").strip()]
    assert texts.count("段落一。") == 1
    assert texts.count("段落二。") == 1


def test_control_children_not_written_twice_at_parent(tmp_path):
    pack = {
        "year": 2026,
        "chapter_name": "设备功能有效性评估",
        "outline": [
            {
                "depth": 1,
                "level": 2,
                "num": "3.4",
                "title": "各线路管控措施",
                "fill": {
                    "empty": False,
                    "source": "管控表",
                    "_children_promoted": True,
                    "sections": [
                        {
                            "title": "轨道交通1号线",
                            "paras": ["加强巡视"],
                            "empty": False,
                            "fill": {"paras": ["加强巡视"], "flow": [], "empty": False},
                        }
                    ],
                },
            },
            {
                "depth": 2,
                "level": 3,
                "num": "3.4.1",
                "title": "轨道交通1号线",
                "content_child": True,
                "fill": {"paras": ["加强巡视"], "flow": [], "empty": False, "source": "管控表"},
            },
        ],
    }
    out = tmp_path / "oh.docx"
    write_chapter_docx(pack, out, "ch3")
    texts = [p.text for p in Document(str(out)).paragraphs if (p.text or "").strip()]
    assert texts.count("加强巡视") == 1


def test_material_diff_not_written_to_word(tmp_path):
    from chapters.overhead.review import review_chapter

    pack = {
        "year": 2026,
        "chapter_name": "设备功能有效性评估",
        "outline": [
            {
                "depth": 1,
                "level": 2,
                "num": "3.3.1.1",
                "title": "各线路柔性接触网状态分布",
                "fill": {
                    "empty": False,
                    "source": "xlsx",
                    "paras": ["表3-5a 各线路柔性接触网状态"],
                    "flow": [{"kind": "para", "text": "表3-5a 各线路柔性接触网状态"}],
                    "table": [["线路", "锚段数量"], ["1号线", "179"]],
                    "warnings": ["5号线在「维护七部A」与「维护七部B」数据不一致，已采用较新文件"],
                },
            }
        ],
    }
    out = tmp_path / "oh.docx"
    write_chapter_docx(pack, out, "ch3")
    texts = [p.text for p in Document(str(out)).paragraphs if (p.text or "").strip()]
    assert not any("材料差异" in t or "数据不一致" in t for t in texts)
    reviews = review_chapter(pack, "ch3")
    assert any(x.get("issue") == "材料差异" and "5号线" in (x.get("note") or "") for x in reviews)


def test_parent_not_yellow_when_child_filled(tmp_path):
    from docx.enum.text import WD_COLOR_INDEX

    pack = {
        "year": 2026,
        "chapter_name": "运营契合满足度评估",
        "outline": [
            {"depth": 0, "level": 1, "num": "4", "title": "运营契合满足度评估"},
            {"depth": 1, "level": 2, "num": "4.1", "title": "设备体量和变化情况"},
            {
                "depth": 2,
                "level": 3,
                "num": "4.1.1",
                "title": "设备体量变化情况",
                "fill": {
                    "empty": False,
                    "source": "七部",
                    "paras": ["18号线新增刚性接触网15.58条公里。"],
                    "flow": [],
                },
            },
        ],
    }
    out = tmp_path / "oh.docx"
    write_chapter_docx(pack, out, "ch4")

    def yellow(p):
        return any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in p.runs)

    heads = {p.text.strip(): yellow(p) for p in Document(str(out)).paragraphs if (p.text or "").strip()}
    parent = next(k for k in heads if k.startswith("4.1") and "设备体量和变化情况" in k)
    child = next(k for k in heads if "设备体量变化情况" in k and k.startswith("4.1.1"))
    assert heads[parent] is False
    assert heads[child] is False
