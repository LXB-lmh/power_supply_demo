# -*- coding: utf-8 -*-
"""插图口径：第1章只留线路概述图、不收故障图；第4章故障图仍写入；第6章不带被丢弃图区。"""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from chapters.power.build import canonical_files
from chapters.power.extract import extract_compliance, extract_overviews, extract_retire
from chapters.power.write import write_chapter_docx
from chapters.ch4.power_extract import build_ch4_pack
from chapters.ch4.power_write import write_ch4_power_docx
from parsers.dispatch import parse_file, parse_files

ROOT = Path(r"F:\材料\2026供电评估_新")
PACK = Path(r"F:\材料\2026供电评估_新\评估材料\99-26年评估报告编制材料")
FAULT = PACK / "评估材料（故障情况，黄色待更新）.docx"
COMP = PACK / "4&5合规性材料(2026).docx"


def _line_report(name_part: str) -> Path | None:
    for p in canonical_files(ROOT):
        if p.suffix.lower() == ".docx" and name_part in p.name:
            return p
    return None


def test_ch1_keeps_overview_maps_not_fault_charts():
    p12 = _line_report("12号线")
    p6 = _line_report("6线路")
    p7 = _line_report("7号线")
    if not (p12 and p6 and p7):
        return
    docs = parse_files([p12, p6, p7])
    by = {x["line"]: x for x in extract_overviews(docs)}
    assert any(x.get("kind") == "drawing" for x in by["12号线"]["flow"])
    assert any("图1" in (x.get("text") or "") for x in by["12号线"]["flow"])
    assert any(x.get("kind") == "drawing" for x in by["6号线"]["flow"])
    assert not any(x.get("kind") == "drawing" for x in (by.get("7号线") or {}).get("flow") or [])  # 7号线概述无线路图，不拿故障图顶上


def test_write_ch1_embeds_line_map(tmp_path):
    p12 = _line_report("12号线")
    if not p12:
        return
    pack = {"year": 2026, "ch1": {"overviews": extract_overviews([parse_file(p12)])}}
    out = tmp_path / "ch1.docx"
    write_chapter_docx(pack, out, "ch1")
    dest = Document(str(out))
    texts = [p.text for p in dest.paragraphs]
    assert any("12号线线路概述" in t for t in texts)
    assert any("图1" in t for t in texts)
    assert dest.element.body.findall(".//" + qn("w:drawing"))


def test_ch4_fault_chart_still_written(tmp_path):
    if not FAULT.exists():
        return
    pack = build_ch4_pack(parse_files([FAULT]), year=2026)
    assert any(x.get("kind") == "drawing" for x in pack["fault_by_line"]["6号线"])
    out = tmp_path / "ch4.docx"
    write_ch4_power_docx(pack, out)
    dest = Document(str(out))
    assert dest.element.body.findall(".//" + qn("w:drawing"))
    xml = dest.element.body.xml
    assert "chart" in xml or "w:drawing" in xml


def test_ch5_ole_and_ch6_without_dropped_zone():
    if not COMP.exists():
        return
    pack = extract_compliance([parse_file(COMP)])
    assert len(pack["ch5"]["drawings"]) >= 3
    assert not any(x.get("kind") == "drawing" for x in pack["ch6"]["intro_flow"])  # 被丢弃图区不进第六章
    assert not any(x.get("kind") == "drawing" for x in pack["ch6"]["revise_flow"])


def test_ch11_still_extracts_sections():
    retire = next((p for p in canonical_files(ROOT) if "11退运" in p.name), None)
    if not retire:
        return
    data = extract_retire([parse_file(retire)])
    assert len(data["sections"]) == 18
    assert any(not s.get("empty") for s in data["sections"])
