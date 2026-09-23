# -*- coding: utf-8 -*-
"""第7章：按表结构与正文语境抽生产计划表（不绑文件名，可逐年复用）。"""
from chapters.power.extract import extract_plan_table, _is_plan_exec_table
from parsers.document_model import Block, DocumentModel


def _doc(name: str, blocks: list[Block], path: str = "") -> DocumentModel:
    return DocumentModel(source_name=name, source_path=path or name, suffix=".docx", blocks=blocks)


def test_plan_table_keeps_only_nearby_caption_and_table():
    plan = [
        ["序号", "线路", "计划数量（项）", "完成数量（项）", "完成率", "备注"],
        ["1", "1号线", "10", "10", "100%", ""],
    ]
    doc = _doc(
        "任意材料名.docx",
        [
            Block(type="paragraph", text="前面无关章节一大段。"),
            Block(type="drawing", source_index=99),
            Block(type="paragraph", text="表7-1生产计划执行情况表"),
            Block(type="table", rows=plan),
            Block(type="drawing", source_index=100),
            Block(type="paragraph", text="后面不该进7.2的运维分析长文。" * 5),
            Block(type="heading", text="8  风险隐患闭环度评估", level=1),
        ],
    )
    hit = extract_plan_table([doc])
    assert hit["table"][1][1] == "1号线"
    assert len(hit["flow"]) <= 2
    assert all(x.get("kind") in {"para", "table"} for x in hit["flow"])
    assert not any(x.get("kind") == "drawing" for x in hit["flow"])
    texts = [x.get("text") or "" for x in hit["flow"] if x.get("kind") == "para"]
    assert any("表7-1" in t for t in texts)
    assert not any("风险隐患" in t for t in texts)


def test_plan_table_by_content_not_filename_rejects_catenary():
    """文件名看起来像供电，正文是接触网状态 → 不进 7.2。"""
    catenary = [
        ["线路", "锚段数量", "A锚段数", "A占比", "B锚段数", "B占比"],
        ["1号线", "179", "26", "14.52%", "33", "18.44%"],
        ["2号线", "100", "10", "10%", "20", "20%"],
        ["3号线", "90", "9", "10%", "18", "20%"],
    ]
    doc = _doc(
        "供电分公司评估材料汇编.docx",
        [
            Block(type="paragraph", text="8.1.1 设备功能有效性"),
            Block(type="paragraph", text="表3-5a 各线路柔性接触网状态"),
            Block(type="table", rows=catenary),
            Block(type="drawing", source_index=12),
            Block(type="paragraph", text="柔性接触网磨耗与导高分析。"),
            Block(type="paragraph", text="刚性接触网拉出值统计。"),
            Block(type="paragraph", text="接触轨磨耗情况。"),
            # 即便夹带一张计划表，全篇接触网信号过强也不该用这份顶 7.2
            Block(
                type="table",
                rows=[
                    ["序号", "线路", "计划数量（项）", "完成数量（项）", "完成率"],
                    ["1", "1号线", "10", "10", "100%"],
                ],
            ),
        ],
    )
    hit = extract_plan_table([doc])
    assert hit["table"] == []


def test_plan_table_accepts_neutral_filename_with_plan_content():
    """文件名无关键字，表头与表题像生产计划 → 应收。"""
    doc = _doc(
        "材料-副本(1).docx",
        [
            Block(type="paragraph", text="日常维修计划执行情况"),
            Block(type="paragraph", text="表7-1生产计划执行情况表"),
            Block(
                type="table",
                rows=[
                    ["序号", "线路", "计划数量（项）", "完成数量（项）", "未完成数量（项）", "完成率", "备注"],
                    ["1", "5号线", "100", "100", "0", "100%", ""],
                    ["2", "6号线", "80", "80", "0", "100%", ""],
                ],
            ),
        ],
    )
    hit = extract_plan_table([doc])
    assert hit["table"][1][1] == "5号线"
    assert hit["source"] == "材料-副本(1).docx"


def test_plan_table_skips_compiled_annual_by_content():
    """多章汇编正文（像完整年报）即使含计划表也不当今年 7.2 唯一来源抢答。"""
    plan = [
        ["序号", "线路", "计划数量（项）", "完成数量（项）", "完成率"],
        ["1", "2号线", "1", "1", "100%"],
    ]
    annual = _doc(
        "材料.docx",
        [
            Block(type="heading", text="3  设备功能有效性评估", level=1),
            Block(type="heading", text="5  管理体系合规性评估", level=1),
            Block(type="heading", text="6  修程修制匹配性评估", level=1),
            Block(type="heading", text="7  运维表现健康度评估", level=1),
            Block(type="heading", text="8  风险隐患闭环度评估", level=1),
            Block(type="table", rows=plan),
        ],
    )
    year = _doc(
        "另一份.docx",
        [
            Block(type="paragraph", text="表7-1生产计划执行情况表"),
            Block(
                type="table",
                rows=[
                    ["序号", "线路", "计划数量（项）", "完成数量（项）", "完成率"],
                    ["1", "3号线", "2", "2", "100%"],
                ],
            ),
        ],
    )
    hit = extract_plan_table([annual, year])
    assert hit["table"][1][1] == "3号线"


def test_plan_exec_layout_keeps_seven_cols_and_header():
    from docx import Document

    from chapters.common.table_copy import apply_plan_exec_layout, plan_exec_widths, write_table_item

    rows = [
        ["序号", "线路", "计划数量（项）", "完成数量（项）", "未完成数量（项）", "完成率", "备注"],
        ["1", "1号线", "1102", "1102", "0", "100%", ""],
        ["2", "2号线", "2083", "2083", "0", "100%", ""],
    ]
    from chapters.common.table_layout import CONTENT_TWIPS

    widths = plan_exec_widths(7)
    assert len(widths) == 7
    assert sum(widths) == CONTENT_TWIPS
    assert widths[5] >= 1000
    doc = Document()
    assert write_table_item(doc, {"kind": "table", "rows": rows, "force_rebuild": True}, domain_id="overhead")
    table = doc.tables[-1]
    apply_plan_exec_layout(table, rows)
    assert len(table.rows) == 3
    assert table.cell(1, 1).text == "1号线"
    assert table.cell(1, 5).text == "100%"


def test_is_plan_exec_table_structure():
    assert not _is_plan_exec_table(
        [["线路", "锚段数量", "A占比", "B占比"], ["1号线", "10", "1%", "2%"]]
    )
    assert _is_plan_exec_table(
        [["序号", "线路", "计划数量（项）", "完成数量（项）", "完成率"], ["1", "1号线", "1", "1", "100%"]]
    )
