# -*- coding: utf-8 -*-
"""上下文逻辑性：A 类与报废同窗应检出。"""
from chapters.common.consistency_logic import check_document_logic


def test_grade_a_vs_scrap_flagged(tmp_path):
    from docx import Document

    path = tmp_path / "bad.docx"
    doc = Document()
    doc.add_heading("3.3 状态分布", level=2)
    doc.add_paragraph("该批整流变压器评估为 A 类设备，运行状态良好。")
    doc.add_paragraph("上述设备绝缘下降严重，无再利用价值，故申请报废。")
    doc.save(path)
    hit = check_document_logic(path, assessment_year=2026, mode="chapter", chapter_id="ch3")
    assert hit["finding_count"] >= 1
    assert any(x.get("issue") == "上下文相悖" for x in hit["findings"])


def test_clean_text_no_false_alarm(tmp_path):
    from docx import Document

    path = tmp_path / "ok.docx"
    doc = Document()
    doc.add_paragraph("1号线变电站温湿度可控，除湿机运行正常。")
    doc.add_paragraph("2号线粉尘区段已采取封堵门缝措施。")
    doc.save(path)
    hit = check_document_logic(path, assessment_year=2026, mode="full")
    assert hit["finding_count"] == 0
