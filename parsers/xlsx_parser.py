# -*- coding: utf-8 -*-
"""Excel：每个工作表一张表。data_only=True 读算出来的值，不读公式文本。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .document_model import Block, DocumentModel


def parse_xlsx(path: Path) -> DocumentModel:
    wb = load_workbook(str(path), data_only=True, read_only=True, keep_links=False)
    try:
        blocks: list[Block] = []
        for sheet in wb.worksheets:
            blocks.append(Block(type="heading", text=f"工作表:{sheet.title}", level=1))
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                values = [("" if c is None else str(c).strip()) for c in row]
                if any(values):
                    rows.append(values)
            if rows:
                blocks.append(Block(type="table", rows=rows))
        return DocumentModel(
            source_name=path.name,
            source_path=str(path),
            suffix=".xlsx",
            blocks=blocks,
        )
    finally:
        wb.close()
